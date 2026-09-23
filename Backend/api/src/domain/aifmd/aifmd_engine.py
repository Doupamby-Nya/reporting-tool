"""AIFMD Annex IV reporting engine (CSSF AIFM Reporting) — full Annex IV coverage.

AIFMD Annex IV reporting is not an EBA regime: it is defined in Level 2
Regulation (EU) No 231/2013 (the AIFMD Implementing Regulation), Annex IV,
with the actual XML schemas published by ESMA and implemented, for
Luxembourg AIFMs, at the CSSF (the national competent authority) under
Circular CSSF 23/833.

Ground truth for every code list, XML element name, nesting path and item
number below is the real ESMA reporting package bundled at
data/regulation/aifmd/esma_package/extracted/2013-1358_.../ :
  - AIFMD_DATMAN_V1.2.xsd / AIFMD_DATAIF_V1.2.xsd / AIFMD_REPORTING_DataTypes_V1.2.xsd
    (the actual XML Schema Definitions — element names, nesting, cardinality)
  - 2013-1358_aifmd_reporting_it_technical_guidance-revision6.xlsx, sheets
    "Technical guidance AIFM file" / "Technical guidance AIF file" (the
    official item-number -> XSD-element-name mapping table used below)
  - Annex II - Excel tables (2013-1360 / 2013-1586): AIFM/AIF reporting
    codes (Table 8/9) and reporting-obligation-change codes (Table 10)
Item numbers (1-38 for the AIFM file, 1-301 for the AIF file) match both
that mapping table and the official "consolidated AIFMD reporting
template v1.1" workbook at data/regulation/aifmd/2013-1359_consolidated_
aifmd_reporting_template-v1.1_revised.xlsx, which Auto mode now serves and
parses directly (see load_official_template_bytes / parse_aifmd_workbook).

Scope: this engine now models every one of the 38 AIFM-file items and 301
AIF-file items (Articles 3(3)(d), 24(1), 24(2) and 24(4)) — identification,
principal markets/instruments, AUM/NAV, investment-strategy breakdown, main
instruments traded, geographical focus, principal exposures and portfolio
concentration, investor concentration, the full Article 24(2) instrument/
turnover/currency exposure tables, dominant-influence companies, market/
counterparty/liquidity risk profiles, financing liquidity, historical
performance, stress-test narratives, rehypothecation, borrowing/leverage,
controlled structures and the five largest borrowing sources. Two areas
remain a labelled best-effort rather than a byte-verified schema fact: (a)
the fine-grained item-by-item conditional-mandatory matrix keyed by AIF
reporting code (the ESMA "Table 9" case list runs to ~120 rows of
AIFM-status/leverage/domicile combinations) is only partially re-implemented
in validation_errors() — the existing high-value checks are kept and a few
new ones added, not the full matrix; and (b) AssetMacroType / AssetType /
market-code free-text guidance defers to the operator's judgement where the
official code still needs a market-specific value (e.g. MIC codes) that no
bundled table enumerates.
"""

from __future__ import annotations

import html as _html
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from domain.aifmd import aifmd_xml as _xml
from domain.shared.xml_common import ERROR, WARNING, Issue, errors_of, warnings_of

REPORTING_MEMBER_STATE = "LU"
SCHEMA_VERSION = "1.2"

from utils.paths import RESOURCES_DIR
OFFICIAL_TEMPLATE_PATH = RESOURCES_DIR / "regulation" / "aifmd" / "2013-1359_consolidated_aifmd_reporting_template-v1.1_revised.xlsx"
TEMPLATE_SHEET_AIFM = "AIFM file 24(1)"
TEMPLATE_SHEET_AIF1 = "AIF file 24(1)"
TEMPLATE_SHEET_AIF2 = "AIF file 24(2)"
TEMPLATE_SHEET_AIF4 = "AIF file 24(4)"

# =============================================================================
# Code lists — every value below is transcribed from a real xs:enumeration in
# AIFMD_REPORTING_DataTypes_V1.2.xsd (or, for the two national/CSSF-only
# lists explicitly marked, from the CSSF AIFM Reporting Technical Guidance
# Chapter 2.2 file-naming convention — not part of the ESMA XSD).
# =============================================================================

FILING_TYPES = {
    "INIT": "Initial reporting for the reporting period",
    "AMND": "Amendment to a reporting previously submitted",
}

# AIFMContentType (item 5, AIFM file) — distinct from AIFMReportingCode (item 16).
AIFM_CONTENT_TYPES = {
    "1": "24(1) reporting contents for all AIFs managed",
    "2": "3(3)(d) reporting contents for all AIFs managed",
    "3": "24(1) reporting contents for all AIFs marketed in the Member State",
}
# AIFContentType (item 5, AIF file) — distinct from AIFReportingCode (item 20).
AIF_CONTENT_TYPES = {
    "1": "24(1) reporting obligation",
    "2": "24(1) + 24(2) reporting obligation",
    "3": "3(3)(d) reporting obligation",
    "4": "24(1) + 24(2) + 24(4) reporting obligation",
    "5": "24(1) + 24(4) reporting obligation",
}

REPORTING_PERIOD_TYPES = ["Q1", "Q2", "Q3", "Q4", "H1", "H2", "Y1", "X1", "X2"]

# Chapter 1.2 of the CSSF Technical Guidance — AIF national code prefix used
# only in the Luxembourg file-naming convention, not an ESMA XSD code list.
AIF_LEGAL_FORM_CODES = {
    "O": "Part II fund or SIF",
    "K": "SICAR",
    "V": "Other vehicle not regulated by the CSSF",
}
# Chapter 2.2 of the CSSF Technical Guidance — sender entity types used only
# in the Luxembourg file-naming convention, not an ESMA XSD code list.
SENDER_ENTITY_TYPES = {
    "B": "Bank", "P": "PFS", "I": "Support PFS", "S": "Management company",
    "K": "Sicar", "1": "Other company not supervised by the CSSF",
}

# AIFMReportingCode (item 16, AIFM file) — Annex II Table 8.
AIFM_REPORTING_CODES = {
    1: "Registered AIFM (Article 3)",
    2: "Authorised AIFM (opt-in) (Article 7)",
    3: "Authorised AIFM with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control (Article 7)",
    4: "Authorised AIFM with half yearly obligation (Article 7)",
    5: "Authorised AIFM with quarterly obligation (Article 7)",
    6: "Non-EU AIFM under Private Placement Regime with annual obligation (Article 42)",
    7: "Non-EU AIFM under Private Placement Regime with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control (Article 42)",
    8: "Non-EU AIFM under Private Placement Regime with half yearly obligation (Article 42)",
    9: "Non-EU AIFM under Private Placement Regime with quarterly obligation (Article 42)",
}

# AIFReportingCode (item 20, AIF file) — Annex II Table 9, transcribed
# verbatim from the AIFReportingCodeType documentation in the real XSD.
AIF_REPORTING_CODES = {
    1: "AIF (of a Registered AIFM)",
    2: "Leveraged EU AIF (of an opt-in AIFM)",
    3: "Leveraged non EU AIF (of an opt-in AIFM) marketed in the Union",
    4: "Leveraged non EU AIF (of an opt-in AIFM), not marketed in the Union",
    5: "Unleveraged EU AIF (of an opt-in AIFM)",
    6: "Unleveraged non EU AIF (of an opt-in AIFM) marketed in the Union",
    7: "Unleveraged non EU AIF (of an opt-in AIFM), not marketed in the Union",
    8: "Unleveraged EU AIF (of an AIFM with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control)",
    9: "Unleveraged non EU AIF (of an AIFM with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control) marketed in the Union",
    10: "Unleveraged non EU AIF (of an AIFM with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control), not marketed in the Union",
    11: "Under 500M Leveraged EU AIF (of an AIFM with half yearly obligation)",
    12: "Under 500M Leveraged non EU AIF (of an AIFM with half yearly obligation) marketed in the Union",
    13: "Under 500M Leveraged non EU AIF (of an AIFM with half yearly obligation), not marketed in the Union",
    14: "Under 500M Unleveraged EU AIF (of an AIFM with half yearly obligation)",
    15: "Under 500M Unleveraged non EU AIF (of an AIFM with half yearly obligation) marketed in the Union",
    16: "Under 500M Unleveraged non EU AIF (of an AIFM with half yearly obligation), not marketed in the Union",
    17: "Unleveraged EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with half yearly obligation)",
    18: "Unleveraged non EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with half yearly obligation) marketed in the Union",
    19: "Unleveraged non EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with half yearly obligation) not marketed in the Union",
    20: "Over 500M Leveraged EU AIF (of an AIFM with half yearly obligation)",
    21: "Over 500M Leveraged non EU AIF (of an AIFM with half yearly obligation) marketed in the Union",
    22: "Over 500M Leveraged non EU AIF (of an AIFM with half yearly obligation), not marketed in the Union",
    23: "Over 500M Unleveraged EU AIF (of an AIFM with half yearly obligation)",
    24: "Over 500M Unleveraged non EU AIF (of an AIFM with half yearly obligation) marketed in the Union",
    25: "Over 500M Unleveraged non EU AIF (of an AIFM with half yearly obligation), not marketed in the Union",
    26: "Unleveraged EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with quarterly obligation)",
    27: "Unleveraged non EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with quarterly obligation) marketed in the Union",
    28: "Unleveraged non EU AIF, investing in non-listed companies in order to acquire control (of an AIFM with quarterly obligation) not marketed in the Union",
    29: "Leveraged EU AIF (of an AIFM with quarterly obligation)",
    30: "Leveraged non EU AIF (of an AIFM with quarterly obligation) marketed in the Union",
    31: "Leveraged non EU AIF (of an AIFM with quarterly obligation), not marketed in the Union",
    32: "Unleveraged EU AIF (of an AIFM with quarterly obligation)",
    33: "Unleveraged non EU AIF (of an AIFM with quarterly obligation) marketed in the Union",
    34: "Unleveraged non EU AIF (of an AIFM with quarterly obligation), not marketed in the Union",
    35: "Leveraged AIF (of an AIFM under Private Placement Regime with yearly obligation) marketed in the Union",
    36: "Unleveraged AIF (of an AIFM under Private Placement Regime with yearly obligation) marketed in the Union",
    37: "Unleveraged AIF (of an AIFM under Private Placement Regime with only unleveraged AIFs investing in non-listed companies and issuers in order to acquire control) marketed in the Union",
    38: "Under 500M Leveraged AIF (of an AIFM under Private Placement Regime with half yearly obligation) marketed in the union",
    39: "Under 500M Unleveraged AIF (of an AIFM under Private Placement Regime with half yearly obligation) marketed in the union",
    40: "Unleveraged AIF, investing in non-listed companies in order to acquire control (of an AIFM under Private Placement Regime with half yearly obligation) marketed in the Union",
    41: "Over 500M Leveraged AIF (of an AIFM under Private Placement Regime with half yearly obligation) marketed in the union",
    42: "Over 500M Unleveraged AIF (of an AIFM under Private Placement Regime with half yearly obligation) marketed in the union",
    43: "Unleveraged AIF, investing in non-listed companies in order to acquire control (of an AIFM under Private Placement Regime with quarterly obligation) marketed in the Union",
    44: "Leveraged AIF (of an AIFM under Private Placement Regime with quarterly obligation) marketed in the Union",
    45: "Unleveraged AIF (of an AIFM under Private Placement Regime with quarterly obligation) marketed in the Union",
}

MASTER_FEEDER_STATUS = {"MASTER": "Master AIF", "FEEDER": "Feeder AIF", "NONE": "Neither master nor feeder"}
FX_EUR_REFERENCE_RATE_TYPES = {"ECB": "ECB-published euro reference rate", "OTH": "Other (free-text source required)"}
MARKET_CODE_TYPES_WITH_NOT = {"MIC": "MIC-coded trading venue", "OTC": "OTC transactions", "XXX": "Not traded on a specific market", "NOT": "No market to report for this rank"}
MARKET_CODE_TYPES = {"MIC": "MIC-coded trading venue", "OTC": "OTC transactions", "XXX": "Not traded on a specific market"}
INSTRUMENT_CODE_TYPES = {"ISIN": "ISIN code", "AII": "Alternative Instrument Identifier (AII)", "NONE": "No code available"}
POSITION_TYPES = {"L": "Long", "S": "Short"}
AII_DERIVATIVE_TYPES = {"O": "Option", "F": "Future/Forward"}
AII_PUT_CALL_TYPES = {"P": "Put", "C": "Call", "F": "Future/Forward (n/a)"}
TYPICAL_POSITION_SIZES = {
    "V_SMALL": "Very small", "SMALL": "Small", "LOW_MID_MKT": "Lower mid market",
    "UP_MID_MKT": "Upper mid market", "L_CAP": "Large cap", "M_CAP": "Mega cap",
}
REDEMPTION_FREQUENCIES = {
    "D": "Daily", "W": "Weekly", "F": "Fortnightly", "M": "Monthly", "Q": "Quarterly",
    "H": "Half-yearly", "Y": "Yearly", "O": "Other", "N": "None",
}
VAR_CALC_METHODS = {"HISTO": "Historical simulation", "CARLO": "Monte Carlo simulation", "PARAM": "Parametric"}
PREDOMINANT_AIF_TYPES = {
    "HFND": "Hedge Fund", "PEQF": "Private Equity Fund", "REST": "Real Estate Fund",
    "FOFS": "Fund of Funds", "OTHR": "Other", "NONE": "None",
}
ASSET_MACRO_TYPES = {
    "SEC": "Securities", "DER": "Derivatives", "CIU": "Collective investment undertakings",
    "PHY": "Physical (real/tangible) assets", "OTH": "Other", "NTA": "No trading activity",
}
TRANSACTION_TYPES = {
    "ACAP": "Acquisition of capital", "BOUT": "Buy-out", "CONS": "Consolidation",
    "CDIV": "Corporate divestment", "ESOP": "Employee share ownership plan", "GCAP": "Growth capital",
    "RCAP": "Replacement capital", "SLIQ": "Secondary liquidity", "TURN": "Turnaround", "OTHR": "Other",
}
INVESTOR_GROUP_TYPES = {
    "NFCO": "Non-financial corporations", "BANK": "Banks",
    "OCIU": "Other collective investment undertaking (e.g. fund of funds or master)",
    "OFIN": "Other financial institutions", "INSC": "Insurance corporations", "PFND": "Pension funds",
    "GENG": "General government", "HHLD": "Households", "UNKN": "Unknown", "NONE": "None",
}

# AssetTypeType (item 104, portfolio concentration — group-level codes), Annex II Table 1.
ASSET_TYPES_CONCENTRATION = {
    "SEC_CSH": "Securities — Cash and cash equivalents", "SEC_LEQ": "Securities — Listed equities",
    "SEC_UEQ": "Securities — Unlisted equities", "SEC_CPN": "Securities — Corporate bonds (non-financial issuer)",
    "SEC_CPI": "Securities — Corporate bonds (financial issuer)", "SEC_SBD": "Securities — Sovereign bonds",
    "SEC_MBN": "Securities — Municipal bonds", "SEC_CBN": "Securities — Convertible bonds (non-financial issuer)",
    "SEC_CBI": "Securities — Convertible bonds (financial issuer)", "SEC_LON": "Securities — Loans",
    "SEC_SSP": "Securities — Structured/securitised products", "DER_EQD": "Derivatives — Equity",
    "DER_FID": "Derivatives — Fixed income", "DER_CDS": "Derivatives — CDS", "DER_FEX": "Derivatives — Foreign exchange",
    "DER_IRD": "Derivatives — Interest rate", "DER_CTY": "Derivatives — Commodity", "DER_OTH": "Derivatives — Other",
    "PHY_RES": "Physical — Real estate", "PHY_CTY": "Physical — Commodities", "PHY_TIM": "Physical — Timber",
    "PHY_ART": "Physical — Art and collectables", "PHY_TPT": "Physical — Transportation assets",
    "PHY_OTH": "Physical — Other", "CIU_OAM": "CIU operated/managed by the AIFM",
    "CIU_NAM": "CIU not operated/managed by the AIFM", "OTH_OTH": "Other asset classes",
    "NTA_NTA": "No trading activity",
}

def _codes(*pairs: tuple[str, str]) -> list[tuple[str, str]]:
    return list(pairs)

# SubAssetTypeType (items 65, 96, 121) — Annex II Table 1, full detail level.
SUB_ASSET_TYPES = _codes(
    ("SEC_CSH_CODP", "Cash and cash equivalents — Certificates of deposit"),
    ("SEC_CSH_COMP", "Cash and cash equivalents — Commercial papers"),
    ("SEC_CSH_OTHD", "Cash and cash equivalents — Other deposits"),
    ("SEC_CSH_OTHC", "Cash and cash equivalents — Other (excl. government securities)"),
    ("SEC_LEQ_IFIN", "Listed equities — Issued by financial institutions"),
    ("SEC_LEQ_OTHR", "Listed equities — Other"),
    ("SEC_UEQ_UEQY", "Unlisted equities"),
    ("SEC_CPN_INVG", "Corporate bonds (non-financial issuer) — Investment grade"),
    ("SEC_CPN_NIVG", "Corporate bonds (non-financial issuer) — Non-investment grade"),
    ("SEC_CPI_INVG", "Corporate bonds (financial issuer) — Investment grade"),
    ("SEC_CPI_NIVG", "Corporate bonds (financial issuer) — Non-investment grade"),
    ("SEC_SBD_EUBY", "Sovereign bonds — EU, 0-1 year to maturity"),
    ("SEC_SBD_EUBM", "Sovereign bonds — EU, 1+ year to maturity"),
    ("SEC_SBD_NOGY", "Sovereign bonds — Non-G10, 0-1 year to maturity"),
    ("SEC_SBD_NOGM", "Sovereign bonds — Non-G10, 1+ year to maturity"),
    ("SEC_SBD_EUGY", "Sovereign bonds — G10 non-EU, 0-1 year to maturity"),
    ("SEC_SBD_EUGM", "Sovereign bonds — G10 non-EU, 1+ year to maturity"),
    ("SEC_MBN_MNPL", "Municipal bonds"),
    ("SEC_CBN_INVG", "Convertible bonds (non-financial issuer) — Investment grade"),
    ("SEC_CBN_NIVG", "Convertible bonds (non-financial issuer) — Non-investment grade"),
    ("SEC_CBI_INVG", "Convertible bonds (financial issuer) — Investment grade"),
    ("SEC_CBI_NIVG", "Convertible bonds (financial issuer) — Non-investment grade"),
    ("SEC_LON_LEVL", "Loans — Leveraged loans"),
    ("SEC_LON_OTHL", "Loans — Other"),
    ("SEC_SSP_SABS", "Structured/securitised products — ABS"),
    ("SEC_SSP_RMBS", "Structured/securitised products — RMBS"),
    ("SEC_SSP_CMBS", "Structured/securitised products — CMBS"),
    ("SEC_SSP_AMBS", "Structured/securitised products — Agency MBS"),
    ("SEC_SSP_ABCP", "Structured/securitised products — ABCP"),
    ("SEC_SSP_CDOC", "Structured/securitised products — CDO/CLO"),
    ("SEC_SSP_STRC", "Structured/securitised products — Structured certificates"),
    ("SEC_SSP_SETP", "Structured/securitised products — ETP"),
    ("SEC_SSP_OTHS", "Structured/securitised products — Other"),
    ("DER_EQD_FINI", "Equity derivatives — Related to financial institutions"),
    ("DER_EQD_OTHD", "Equity derivatives — Other"),
    ("DER_FID_FIXI", "Fixed income derivatives"),
    ("DER_CDS_SNFI", "CDS — Single name financial"),
    ("DER_CDS_SNSO", "CDS — Single name sovereign"),
    ("DER_CDS_SNOT", "CDS — Single name other"),
    ("DER_CDS_INDX", "CDS — Index"),
    ("DER_CDS_EXOT", "CDS — Exotic (incl. credit default tranche)"),
    ("DER_CDS_OTHR", "CDS — Other"),
    ("DER_FEX_INVT", "Foreign exchange — For investment purposes"),
    ("DER_FEX_HEDG", "Foreign exchange — For hedging purposes"),
    ("DER_IRD_INTR", "Interest rate derivatives"),
    ("DER_CTY_ECOL", "Commodity derivatives — Crude oil"),
    ("DER_CTY_ENNG", "Commodity derivatives — Natural gas"),
    ("DER_CTY_ENPW", "Commodity derivatives — Power"),
    ("DER_CTY_ENOT", "Commodity derivatives — Other energy"),
    ("DER_CTY_PMGD", "Commodity derivatives — Gold"),
    ("DER_CTY_PMOT", "Commodity derivatives — Other precious metals"),
    ("DER_CTY_OTIM", "Commodity derivatives — Industrial metals"),
    ("DER_CTY_OTLS", "Commodity derivatives — Livestock"),
    ("DER_CTY_OTAP", "Commodity derivatives — Agricultural products"),
    ("DER_CTY_OTHR", "Commodity derivatives — Other commodities"),
    ("DER_OTH_OTHR", "Other derivatives"),
    ("PHY_RES_RESL", "Physical — Residential real estate"),
    ("PHY_RES_COML", "Physical — Commercial real estate"),
    ("PHY_RES_OTHR", "Physical — Other real estate"),
    ("PHY_CTY_PCTY", "Physical — Commodities"),
    ("PHY_TIM_PTIM", "Physical — Timber"),
    ("PHY_ART_PART", "Physical — Art and collectables"),
    ("PHY_TPT_PTPT", "Physical — Transportation assets"),
    ("PHY_OTH_OTHR", "Physical — Other"),
    ("CIU_OAM_MMFC", "CIU operated/managed by the AIFM — Money Market/Cash management CIU"),
    ("CIU_OAM_AETF", "CIU operated/managed by the AIFM — ETF"),
    ("CIU_OAM_OTHR", "CIU operated/managed by the AIFM — Other CIU"),
    ("CIU_NAM_MMFC", "CIU not operated/managed by the AIFM — Money Market/Cash management CIU"),
    ("CIU_NAM_AETF", "CIU not operated/managed by the AIFM — ETF"),
    ("CIU_NAM_OTHR", "CIU not operated/managed by the AIFM — Other CIU"),
    ("OTH_OTH_OTHR", "Other asset classes"),
    ("NTA_NTA_NOTA", "No trading activity in this sub-asset type"),
)
SUB_ASSET_TYPE_LABELS = dict(SUB_ASSET_TYPES)

# TurnoverSubAssetTypeType (items 125-127) — Annex II Table 1, turnover-level detail.
TURNOVER_SUB_ASSET_TYPES = _codes(
    ("SEC_CSH_CSH", "Cash and cash equivalents"), ("SEC_LEQ_LEQ", "Listed equities"),
    ("SEC_UEQ_UEQ", "Unlisted equities"), ("SEC_CPN_IVG", "Corporate bonds (non-FI) — Investment grade"),
    ("SEC_CPN_NIG", "Corporate bonds (non-FI) — Non-investment grade"),
    ("SEC_CPI_CPI", "Corporate bonds (financial institutions)"),
    ("SEC_SBD_EUB", "Sovereign bonds — EU Member State"), ("SEC_SBD_NEU", "Sovereign bonds — Non-EU Member State"),
    ("SEC_MUN_MUN", "Municipal bonds"), ("SEC_CBD_CBD", "Convertible bonds"), ("SEC_LON_LON", "Loans"),
    ("SEC_SSP_SSP", "Structured/securitised products"), ("DER_EQD_EQD", "Equity derivatives"),
    ("DER_FID_FID", "Fixed income derivatives"), ("DER_CDS_CDS", "CDS"),
    ("DER_FEX_INV", "Foreign exchange — For investment purposes"),
    ("DER_FEX_HED", "Foreign exchange — For hedging purposes"), ("DER_IRD_IRD", "Interest rate derivatives"),
    ("DER_CTY_CTY", "Commodity derivatives"), ("DER_OTH_OTH", "Other derivatives"),
    ("PHY_RES_RES", "Physical — Real estate"), ("PHY_CTY_CTY", "Physical — Commodities"),
    ("PHY_TIM_TIM", "Physical — Timber"), ("PHY_ART_ART", "Physical — Art and collectables"),
    ("PHY_TPT_TPT", "Physical — Transportation assets"), ("PHY_OTH_OTH", "Physical — Other"),
    ("CIU_CIU_CIU", "Collective investment undertakings"), ("OTH_OTH_OTH", "Other asset classes"),
)
TURNOVER_SUB_ASSET_TYPE_LABELS = dict(TURNOVER_SUB_ASSET_TYPES)

# HedgeFundStrategyTypeType / PrivateEquityFundStrategyTypeType / RealEstateFundStrategyTypeType /
# FundOfFundsStrategyTypeType / OtherFundStrategyTypeType (items 58-61, one enum per predominant AIF type).
HEDGE_FUND_STRATEGIES = _codes(
    ("EQTY_LGBS", "Equity: Long Bias"), ("EQTY_LGST", "Equity: Long/Short"),
    ("EQTY_MTNL", "Equity: Market Neutral"), ("EQTY_STBS", "Equity: Short Bias"),
    ("RELV_FXIA", "Relative Value: Fixed Income Arbitrage"), ("RELV_CBAR", "Relative Value: Convertible Bond Arbitrage"),
    ("RELV_VLAR", "Relative Value: Volatility Arbitrage"), ("EVDR_DSRS", "Event Driven: Distressed/Restructuring"),
    ("EVDR_RAMA", "Event Driven: Risk Arbitrage/Merger Arbitrage"), ("EVDR_EYSS", "Event Driven: Equity Special Situations"),
    ("CRED_LGST", "Credit Long/Short"), ("CRED_ABLG", "Credit Asset Based Lending"), ("MACR_MACR", "Macro"),
    ("MANF_CTAF", "Managed Futures/CTA: Fundamental"), ("MANF_CTAQ", "Managed Futures/CTA: Quantitative"),
    ("MULT_HFND", "Multi-strategy hedge fund"), ("OTHR_HFND", "Other hedge fund strategy"),
)
PRIVATE_EQUITY_STRATEGIES = _codes(
    ("VENT_CAPL", "Venture Capital"), ("GRTH_CAPL", "Growth Capital"), ("MZNE_CAPL", "Mezzanine Capital"),
    ("MULT_PEQF", "Multi-strategy private equity fund"), ("OTHR_PEQF", "Other private equity fund strategy"),
)
REAL_ESTATE_STRATEGIES = _codes(
    ("RESL_REST", "Residential real estate"), ("COML_REST", "Commercial real estate"),
    ("INDL_REST", "Industrial real estate"), ("MULT_REST", "Multi-strategy real estate fund"),
    ("OTHR_REST", "Other real estate strategy"),
)
FUND_OF_FUNDS_STRATEGIES = _codes(
    ("FOFS_FHFS", "Fund of hedge funds"), ("FOFS_PRIV", "Fund of private equity"), ("OTHR_FOFS", "Other fund of funds"),
)
OTHER_STRATEGIES = _codes(
    ("OTHR_COMF", "Commodity fund"), ("OTHR_EQYF", "Equity fund"), ("OTHR_FXIF", "Fixed income fund"),
    ("OTHR_INFF", "Infrastructure fund"), ("OTHR_OTHF", "Other fund"),
)
STRATEGY_GROUPS: list[tuple[str, str, list[tuple[str, str]]]] = [
    ("HFND", "a) Hedge Fund Strategies", HEDGE_FUND_STRATEGIES),
    ("PEQF", "b) Private Equity Strategies", PRIVATE_EQUITY_STRATEGIES),
    ("REST", "c) Real Estate Strategies", REAL_ESTATE_STRATEGIES),
    ("FOFS", "d) Fund of Fund Strategies", FUND_OF_FUNDS_STRATEGIES),
    ("OTHR", "e) Other Strategies", OTHER_STRATEGIES),
]

# NAVGeographicalFocus / AUMGeographicalFocus (items 78-93) — 8 fixed regions, each its own scalar element.
GEO_REGIONS: list[tuple[str, str, str]] = [
    ("AFRICA", "AfricaNAVRate", "AfricaAUMRate"), ("ASIA_PACIFIC", "AsiaPacificNAVRate", "AsiaPacificAUMRate"),
    ("EUROPE_OTHER", "EuropeNAVRate", "EuropeAUMRate"), ("EEA", "EEANAVRate", "EEAAUMRate"),
    ("MIDDLE_EAST", "MiddleEastNAVRate", "MiddleEastAUMRate"), ("NORTH_AMERICA", "NorthAmericaNAVRate", "NorthAmericaAUMRate"),
    ("SOUTH_AMERICA", "SouthAmericaNAVRate", "SouthAmericaAUMRate"), ("SUPRANATIONAL", "SupraNationalNAVRate", "SupraNationalAUMRate"),
]
GEO_REGION_LABELS = {
    "AFRICA": "Africa", "ASIA_PACIFIC": "Asia and Pacific (other than Middle East)",
    "EUROPE_OTHER": "Europe (other than EEA)", "EEA": "Europe (EEA)", "MIDDLE_EAST": "Middle East",
    "NORTH_AMERICA": "North America", "SOUTH_AMERICA": "South America", "SUPRANATIONAL": "Supranational / multiple region",
}

# Portfolio/Investor liquidity + financing-maturity buckets (items 178-185, 186-192, 210-217) — same 7 day-buckets.
LIQUIDITY_BUCKETS: list[tuple[str, str]] = [
    ("D0_1", "1 day or less"), ("D2_7", "2 - 7 days"), ("D8_30", "8 - 30 days"), ("D31_90", "31 - 90 days"),
    ("D91_180", "91 - 180 days"), ("D181_365", "181 - 365 days"), ("D365_MORE", "More than 365 days"),
]

# MarketRiskMeasure codes (items 138-147) — real ESMA/CSSF risk-measure codes;
# the first three are mandatory whenever the AIF is subject to Article 24(2) reporting (CAF-216).
MARKET_RISK_MEASURE_TYPES: list[tuple[str, str, bool]] = [
    ("NET_EQTY_DELTA", "Net Equity Delta", True),
    ("NET_CS01", "Net CS01", True),
    ("NET_DV01", "Net DV01", True),
    ("VAR", "VAR", False),
    ("NET_FX_DELTA", "Net FX Delta", False),
    ("NET_CTY_DELTA", "Net Commodity Delta", False),
    ("VEGA_EXPO", "Vega Exposure", False),
]
MANDATORY_RISK_MEASURE_TYPES = [code for code, _, mandatory in MARKET_RISK_MEASURE_TYPES if mandatory]

# CAF-206 / CAF-207 — AIFReportingCode values for which the liquidity fields
# and the mandatory market risk measures (CAF-216) are required under Article 24(2).
ARTICLE_24_2_CODES = {
    2, 3, 5, 6, 8, 9, 11, 12, 14, 15, 17, 18, 20, 21, 23, 24, 26, 27, 29, 30,
    32, 33, 35, 36, 37, 38, 39, 40, 41, 42, 43, 44, 45,
}
# CAF-208 / CAF-209 / CAF-210 / CAF-211 / CAF-212 — AIFReportingCode values
# for which leverage/borrowing fields are mandatory under Article 24(2) AND 24(4).
ARTICLE_24_2_AND_24_4_CODES = {
    2, 3, 4, 5, 6, 8, 9, 11, 12, 13, 14, 15, 17, 18, 20, 21, 22, 23, 24, 26,
    27, 29, 30, 31, 32, 33, 35, 36, 37, 38, 39, 40, 41, 42, 43, 44, 45,
}

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def _esc(v: Any) -> str:
    return _html.escape(str(v))


def _tag(name: str, value: Any) -> str:
    if value is None or value == "":
        return ""
    return f"<{name}>{_esc(value)}</{name}>"


def _bool_str(v: bool | None) -> str | None:
    if v is None:
        return None
    return "true" if v else "false"


def _btag(name: str, value: bool | None) -> str:
    return _tag(name, _bool_str(value))


def _rows_xml(rows: list[dict], row_tag: str, columns: list[tuple[str, str]]) -> str:
    """Generic serializer for a repeating group: ``columns`` is a list of
    (dict_key, xml_tag) pairs, in schema document order."""
    out = []
    for row in rows:
        inner = "".join(_tag(tag, row.get(key)) for key, tag in columns)
        out.append(f"<{row_tag}>{inner}</{row_tag}>")
    return "".join(out)


# =============================================================================
# Data model
# =============================================================================

@dataclass
class AifmdSubmission:
    # --- Sender identification (CSSF Chapter 2.2 file-naming convention) ---
    sender_entity_type: str = "S"
    sender_id: str = ""

    # Reporting member state of the file (CSSF control FIL-009 requires "LU"
    # for a Luxembourg filing); carried as data so a test or a non-LU filing
    # can set it explicitly rather than relying on a module constant.
    reporting_member_state: str = REPORTING_MEMBER_STATE

    # A cancellation file carries a CancellationAIF(M)RecordInfo record
    # instead of the report itself (CSSF controls CSSF-009, CSSF-013).
    aifm_cancellation: bool = False
    aif_cancellation: bool = False

    # ---------------------------------------------------------------
    # AIFM file — items 1-38 (AIFMReportingInfo/AIFMRecordInfo)
    # ---------------------------------------------------------------
    aifm_version: str = SCHEMA_VERSION                       # item 2
    aifm_filing_type: str = "INIT"                            # item 4
    aifm_content_type: str = "1"                               # item 5
    reporting_period_type: str = "Y1"                          # item 8 (shared AIFM/AIF)
    reporting_period_year: int = field(default_factory=lambda: date.today().year - 1)  # item 9
    reporting_period_start: date = field(default_factory=lambda: date(date.today().year - 1, 1, 1))  # item 6
    reporting_period_end: date = field(default_factory=lambda: date(date.today().year - 1, 12, 31))  # item 7
    aifm_change_frequency_code: str | None = None              # item 10
    aifm_change_contents_code: str | None = None                # item 11
    aifm_change_quarter: str | None = None                      # item 12
    aifm_last_reporting_flag: bool = False                      # item 13
    aifm_assumptions: list[dict] = field(default_factory=list)  # items 14-15
    aifm_reporting_code: int = 1                                 # item 16
    aifm_jurisdiction: str = "LU"                                # item 17
    aifm_national_code: str = ""                                  # item 18, "A" + 8 digits
    aifm_name: str = ""                                            # item 19
    aifm_eea_flag: bool = True                                     # item 20
    aifm_no_reporting_flag: bool = False                           # item 21
    aifm_lei: str = ""                                              # item 22
    aifm_bic: str = ""                                               # item 23
    aifm_old_identifier_country: str = ""                            # item 24
    aifm_old_identifier_code: str = ""                                # item 25
    aifm_principal_markets: list[dict] = field(default_factory=list)  # items 26-29 (5 ranked rows)
    aifm_principal_instruments: list[dict] = field(default_factory=list)  # items 30-32 (5 ranked rows)
    total_aum_eur: float = 0.0                                          # item 33
    total_aum_base_currency: float = 0.0                                 # item 34
    aifm_base_currency: str = "EUR"                                       # item 35
    aifm_fx_ref_rate_type: str = ""                                        # item 36
    eur_fx_rate: float = 1.0                                                # item 37
    aifm_fx_ref_rate_description: str = ""                                  # item 38

    # ---------------------------------------------------------------
    # AIF file — Header + Article 24(1) — items 1-120
    # ---------------------------------------------------------------
    aif_version: str = SCHEMA_VERSION                            # item 2
    aif_filing_type: str = "INIT"                                 # item 4
    aif_content_type: str = "2"                                    # item 5
    aif_change_frequency_code: str | None = None                    # item 10
    aif_change_contents_code: str | None = None                      # item 11
    aif_change_quarter: str | None = None                             # item 12
    aif_last_reporting_flag: bool = False                              # item 13
    aif_assumptions: list[dict] = field(default_factory=list)          # items 14-15

    aif_legal_form: str = "O"                          # CSSF file-naming prefix ('O' Part II/SIF, 'K' Sicar, 'V' other)
    aif_national_number: str = ""                        # CSSF file-naming, 8 digits
    aif_subfund_number: str = "00000000"                   # CSSF file-naming, 8 digits
    aif_name: str = ""                                       # item 18
    aif_eea_flag: bool = True                                 # item 19
    aif_reporting_code: int = 1                                 # item 20
    aif_domicile: str = ""                                        # item 21
    aif_inception_date: date | None = None                         # item 22
    aif_no_reporting_flag: bool = False                              # item 23

    aif_lei: str = ""                    # item 24
    aif_isin: str = ""                    # item 25
    aif_cusip: str = ""                    # item 26
    aif_sedol: str = ""                     # item 27
    aif_bloomberg: str = ""                  # item 28
    aif_reuters: str = ""                     # item 29
    aif_ecb_code: str = ""                     # item 30
    aif_old_identifier_country: str = ""        # item 31
    aif_old_identifier_code: str = ""             # item 32

    aif_share_class_flag: bool = False              # item 33
    aif_share_classes: list[dict] = field(default_factory=list)  # items 34-40

    aif_master_feeder_status: str = "NONE"           # item 41
    aif_master_aifs: list[dict] = field(default_factory=list)     # items 42-44
    aif_prime_brokers: list[dict] = field(default_factory=list)    # items 45-47

    aif_aum_base_currency: float = 0.0                # item 48
    aif_base_currency: str = "EUR"                      # item 49
    aif_fx_rate: float = 1.0                              # item 50
    aif_fx_ref_rate_type: str = ""                          # item 51
    aif_fx_ref_rate_description: str = ""                    # item 52
    aif_nav: float = 0.0                                       # item 53

    funding_source_country_1: str = ""     # item 54
    funding_source_country_2: str = ""      # item 55
    funding_source_country_3: str = ""       # item 56

    predominant_aif_type: str = "NONE"                  # item 57
    investment_strategies: list[dict] = field(default_factory=list)  # items 58-61

    hft_transaction_number: int | None = None          # item 62
    hft_buy_sell_market_value: float | None = None       # item 63

    main_instruments: list[dict] = field(default_factory=list)  # items 64-77 (5 ranked rows)
    nav_geo_focus: dict[str, float] = field(default_factory=dict)   # items 78-85
    aum_geo_focus: dict[str, float] = field(default_factory=dict)    # items 86-93
    ten_principal_exposures: list[dict] = field(default_factory=list)  # items 94-102 (10 ranked rows)
    portfolio_concentrations: list[dict] = field(default_factory=list)  # items 103-112 (5 ranked rows)
    typical_position_size: str = ""                       # item 113
    aif_principal_markets: list[dict] = field(default_factory=list)  # items 114-117 (3 ranked rows)
    investor_beneficial_owners_rate: float | None = None  # item 118
    investor_professional_rate: float | None = None         # item 119
    investor_retail_rate: float | None = None                 # item 120

    # ---------------------------------------------------------------
    # AIF file — Article 24(2) — items 121-280
    # ---------------------------------------------------------------
    instrument_exposures: list[dict] = field(default_factory=list)  # items 121-124
    asset_turnovers: list[dict] = field(default_factory=list)         # items 125-127
    currency_exposures: list[dict] = field(default_factory=list)        # items 128-130
    dominant_influence_companies: list[dict] = field(default_factory=list)  # items 131-136

    annual_investment_return_rate: str = ""              # item 137 ("NA" or a signed %)
    market_risk_measures: list[dict] = field(default_factory=list)  # items 138-147

    trading_pct_securities_exchange: float | None = None  # item 148
    trading_pct_securities_otc: float | None = None          # item 149
    trading_pct_derivatives_exchange: float | None = None      # item 150
    trading_pct_derivatives_otc: float | None = None             # item 151
    derivatives_ccp_rate: float | None = None                       # item 152
    derivatives_bilateral_rate: float | None = None                    # item 153
    repo_ccp_rate: float | None = None                                    # item 154
    repo_bilateral_rate: float | None = None                                # item 155
    repo_triparty_rate: float | None = None                                   # item 156

    collateral_cash_posted: float | None = None            # item 157
    collateral_securities_posted: float | None = None         # item 158
    collateral_other_posted: float | None = None                # item 159

    counterparty_exposures_to_aif: list[dict] = field(default_factory=list)   # items 160-165 (5 rows)
    counterparty_exposures_from_aif: list[dict] = field(default_factory=list)  # items 166-171 (5 rows)
    ccp_clearing_flag: bool | None = None                    # item 172
    ccp_exposures: list[dict] = field(default_factory=list)   # items 173-177 (up to 3 rows)

    portfolio_liquidity: dict[str, float] = field(default_factory=dict)  # items 178-184
    unencumbered_cash: float | None = None                                 # item 185
    investor_liquidity: dict[str, float] = field(default_factory=dict)      # items 186-192

    provide_withdrawal_rights: bool | None = None       # item 193
    redemption_frequency: str = ""                         # item 194
    redemption_notice_period_days: int | None = None          # item 195
    redemption_lockup_period_days: int | None = None            # item 196

    side_pocket_rate: float | None = None            # item 197
    gates_rate: float | None = None                    # item 198
    dealing_suspension_rate: float | None = None          # item 199
    other_arrangement_type: str = ""                         # item 200
    other_arrangement_rate: float | None = None                 # item 201
    total_arrangement_rate: float | None = None                   # item 202
    investor_preferential_treatment_flag: bool | None = None        # item 203
    pref_disclosure_flag: bool | None = None                          # item 204
    pref_liquidity_flag: bool | None = None                             # item 205
    pref_fee_flag: bool | None = None                                     # item 206
    pref_other_flag: bool | None = None                                     # item 207

    investor_groups: list[dict] = field(default_factory=list)  # items 208-209 (10 fixed rows)

    total_financing_amount: float | None = None                # item 210
    financing_buckets: dict[str, float] = field(default_factory=dict)  # items 211-217

    total_open_positions: int | None = None    # item 218
    historical_performance: list[dict] = field(default_factory=list)  # items 219-278 (12 rows)

    stress_test_article15: str = ""     # item 279
    stress_test_article16: str = ""       # item 280

    # ---------------------------------------------------------------
    # AIF file — Article 24(2)+(4) shared — items 281-295
    # ---------------------------------------------------------------
    rehypothecation_flag: bool | None = None       # item 281
    rehypothecation_rate: float | None = None        # item 282
    unsecured_borrowing_amount: float | None = None    # item 283
    secured_borrowing_prime_brokerage_amount: float | None = None  # item 284
    secured_borrowing_reverse_repo_amount: float | None = None       # item 285
    secured_borrowing_other_amount: float | None = None                # item 286
    exchange_traded_derivatives_exposure: float | None = None            # item 287
    otc_derivatives_amount: float | None = None                            # item 288
    short_position_borrowed_securities_value: float | None = None            # item 289
    controlled_structures: list[dict] = field(default_factory=list)            # items 290-293
    gross_method_rate: float | None = None            # item 294
    commitment_method_rate: float | None = None         # item 295

    # ---------------------------------------------------------------
    # AIF file — Article 24(4) only — items 296-301
    # ---------------------------------------------------------------
    borrowing_sources: list[dict] = field(default_factory=list)  # items 296-301 (5 fixed rows)

    @property
    def total_aum_eur_computed(self) -> float:
        """CSSF-031: if the base currency is EUR this must equal the base-
        currency amount; otherwise base-currency amount x EUR FX rate."""
        if (self.aifm_base_currency or "").upper() == "EUR":
            return self.total_aum_base_currency
        return self.total_aum_base_currency * self.eur_fx_rate

    @property
    def aif_national_code_full(self) -> str:
        """AIF file item 17 — CSSF derives this from the Chapter 2.2.2
        file-naming components (legal-form prefix + 8-digit number)."""
        return f"{self.aif_legal_form}{self.aif_national_number}"

    def validation_issues(self) -> list["Issue"]:
        """Every CSSF control that applies to this submission, both files, each
        carrying the control's own code (see domain.aifmd.aifmd_validation)."""
        from domain.aifmd import aifmd_validation
        return aifmd_validation.validation_issues(self)

    def validation_errors(self) -> list[str]:
        """The blocking issues only, as plain strings (kept for callers that
        predate validation_issues)."""
        return [str(i) for i in errors_of(self.validation_issues())]


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d%H%M%S")


def _now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%dT%H:%M:%S")


def build_datman_filename(sub: AifmdSubmission) -> str:
    """AIFM-level file name — Chapter 2.2.1 of the CSSF Technical Guidance:
    AIFREP-<sender type><sender id>-A<AIFM id>-F00000000-00000000-MAN-<timestamp>.xml"""
    return (
        f"AIFREP-{sub.sender_entity_type}{(sub.sender_id or '0').zfill(8)}-"
        f"{sub.aifm_national_code}-F00000000-00000000-MAN-{_timestamp()}.xml"
    )


def build_dataif_filename(sub: AifmdSubmission) -> str:
    """AIF-level file name — Chapter 2.2.2 of the CSSF Technical Guidance:
    AIFREP-<sender type><sender id>-A<AIFM id>-<O|K|V><AIF id>-<subfund id>-AIF-<timestamp>.xml"""
    return (
        f"AIFREP-{sub.sender_entity_type}{(sub.sender_id or '0').zfill(8)}-"
        f"{sub.aifm_national_code}-{sub.aif_legal_form}{sub.aif_national_number}-"
        f"{sub.aif_subfund_number}-AIF-{_timestamp()}.xml"
    )


def build_datman_xml(sub: AifmdSubmission, creation: str | None = None) -> str:
    """AIFM-level reporting file (DATMAN). The XML shape lives in
    domain.aifmd.aifmd_xml, which is built against the real ESMA schema and is
    validated against it by domain.aifmd.aifmd_xsd."""
    return _xml.build_datman_xml(sub, creation)


def build_dataif_xml(sub: AifmdSubmission, creation: str | None = None) -> str:
    """AIF-level reporting file (DATAIF) — see build_datman_xml."""
    return _xml.build_dataif_xml(sub, creation)


def bucket_or_vega_present(r: dict) -> bool:
    return any(r.get(k) not in (None, "") for k in ("bucket_lt5", "bucket_5_15", "bucket_gt15", "vega_current", "vega_lower10", "vega_higher10", "var_value", "var_method", "description"))


# =============================================================================
# Auto mode — the official ESMA consolidated reporting template, used as-is.
# =============================================================================

def load_official_template_bytes() -> bytes:
    """Returns the exact bytes of the official ESMA consolidated AIFMD
    reporting template v1.1 (data/regulation/aifmd/2013-1359_...xlsx) — the
    same file published at
    https://www.esma.europa.eu/document/consolidated-aifmd-reporting-template-revised,
    served as-is (no re-generation) so a preparer can fill it in Excel using
    the ESMA-authored 'Data Type' / 'Reported Data' layout directly."""
    return OFFICIAL_TEMPLATE_PATH.read_bytes()


def _cell(ws, row: int, col: str):
    return ws[f"{col}{row}"].value


def _clean(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _num(v: Any) -> float | None:
    v = _clean(v)
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _int(v: Any) -> int | None:
    n = _num(v)
    return int(n) if n is not None else None


def parse_aifmd_workbook(file_bytes: bytes) -> tuple[dict[str, Any], dict[str, Any], list[str]]:
    """Parses the official ESMA consolidated reporting template (as served
    by load_official_template_bytes(), filled in by a preparer) back into
    (aifm_info dict, aif_info dict, errors) — the two dicts use the same
    snake_case keys as AifmdSubmission's constructor args (see
    submission_from_dicts) and feed the Manual-mode form 1:1, so a workbook
    round-trips into the exact same widgets a preparer would otherwise fill
    in by hand. Simple (single-cell) items are read generically from column
    A's item number; the ~35 repeating/table blocks are read from their
    known column layout (row/column positions confirmed against the real
    template — see the module docstring)."""
    import openpyxl

    errors: list[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        return {}, {}, [f"Could not read this workbook: {exc}"]

    for sheet in (TEMPLATE_SHEET_AIFM, TEMPLATE_SHEET_AIF1, TEMPLATE_SHEET_AIF2, TEMPLATE_SHEET_AIF4):
        if sheet not in wb.sheetnames:
            errors.append(f"Sheet '{sheet}' not found — is this the official ESMA consolidated template?")

    aifm: dict[str, Any] = {}
    aif: dict[str, Any] = {}

    def simple(ws, item_to_key: dict[int, str], value_col: str = "C", max_row: int = 400) -> dict[str, Any]:
        out: dict[str, Any] = {}
        for r in range(1, max_row + 1):
            item_no = _cell(ws, r, "A")
            if isinstance(item_no, (int, float)) and int(item_no) in item_to_key:
                out[item_to_key[int(item_no)]] = _clean(_cell(ws, r, value_col))
        return out

    if TEMPLATE_SHEET_AIFM in wb.sheetnames:
        ws = wb[TEMPLATE_SHEET_AIFM]
        aifm.update(simple(ws, {
            2: "aifm_version", 4: "aifm_filing_type", 5: "aifm_content_type",
            6: "reporting_period_start", 7: "reporting_period_end", 8: "reporting_period_type", 9: "reporting_period_year",
            10: "aifm_change_frequency_code", 11: "aifm_change_contents_code", 12: "aifm_change_quarter",
            13: "aifm_last_reporting_flag", 16: "aifm_reporting_code", 17: "aifm_jurisdiction",
            18: "aifm_national_code", 19: "aifm_name", 20: "aifm_eea_flag", 21: "aifm_no_reporting_flag",
            22: "aifm_lei", 23: "aifm_bic", 24: "aifm_old_identifier_country", 25: "aifm_old_identifier_code",
            33: "total_aum_eur", 34: "total_aum_base_currency", 35: "aifm_base_currency",
            36: "aifm_fx_ref_rate_type", 37: "eur_fx_rate", 38: "aifm_fx_ref_rate_description",
        }))
        markets = []
        for r in range(42, 47):
            markets.append({"ranking": _cell(ws, r, "C"), "market_code_type": _clean(_cell(ws, r, "D")), "market_code": _clean(_cell(ws, r, "E")), "aggregated_value": _num(_cell(ws, r, "F"))})
        aifm["aifm_principal_markets"] = [m for m in markets if m.get("market_code_type") or m.get("aggregated_value")]
        instruments = []
        for r in range(50, 55):
            instruments.append({"ranking": r - 49, "sub_asset_type": _clean(_cell(ws, r, "C")), "aggregated_value": _num(_cell(ws, r, "E"))})
        aifm["aifm_principal_instruments"] = [i for i in instruments if i.get("sub_asset_type") or i.get("aggregated_value")]

    if TEMPLATE_SHEET_AIF1 in wb.sheetnames:
        ws = wb[TEMPLATE_SHEET_AIF1]
        aif.update(simple(ws, {
            2: "aif_version", 4: "aif_filing_type", 5: "aif_content_type",
            10: "aif_change_frequency_code", 11: "aif_change_contents_code", 12: "aif_change_quarter", 13: "aif_last_reporting_flag",
            17: "aif_national_code_full", 18: "aif_name", 19: "aif_eea_flag", 20: "aif_reporting_code",
            21: "aif_domicile", 22: "aif_inception_date", 23: "aif_no_reporting_flag",
            24: "aif_lei", 25: "aif_isin", 26: "aif_cusip", 27: "aif_sedol", 28: "aif_bloomberg", 29: "aif_reuters", 30: "aif_ecb_code",
            31: "aif_old_identifier_country", 32: "aif_old_identifier_code", 33: "aif_share_class_flag",
            41: "aif_master_feeder_status", 48: "aif_aum_base_currency", 49: "aif_base_currency", 50: "aif_fx_rate",
            51: "aif_fx_ref_rate_type", 52: "aif_fx_ref_rate_description", 53: "aif_nav",
            54: "funding_source_country_1", 55: "funding_source_country_2", 56: "funding_source_country_3",
            57: "predominant_aif_type", 62: "hft_transaction_number", 63: "hft_buy_sell_market_value",
            113: "typical_position_size", 118: "investor_beneficial_owners_rate",
            119: "investor_professional_rate", 120: "investor_retail_rate",
        }, max_row=203))
        # 8-region geographical focus tables — fixed rows, value in column C.
        nav_geo = {}
        for r, key in zip(range(149, 157), [k for k, _, _ in GEO_REGIONS]):
            nav_geo[key] = _num(_cell(ws, r, "C"))
        aif["nav_geo_focus"] = nav_geo
        aum_geo = {}
        for r, key in zip(range(159, 167), [k for k, _, _ in GEO_REGIONS]):
            aum_geo[key] = _num(_cell(ws, r, "C"))
        aif["aum_geo_focus"] = aum_geo
        # 10 principal exposures (rows 170-179), 5 portfolio concentrations (183-187), 3 AIF principal markets (194-196).
        exposures = []
        for r in range(170, 180):
            exposures.append({
                "ranking": _cell(ws, r, "C"), "macro_type": _clean(_cell(ws, r, "D")), "sub_asset_type": _clean(_cell(ws, r, "F")),
                "position_type": _clean(_cell(ws, r, "G")), "aggregated_value": _num(_cell(ws, r, "H")), "aggregated_value_pct": _num(_cell(ws, r, "I")),
                "counterparty_name": _clean(_cell(ws, r, "J")), "counterparty_lei": _clean(_cell(ws, r, "K")), "counterparty_bic": _clean(_cell(ws, r, "L")),
            })
        aif["ten_principal_exposures"] = [e for e in exposures if e.get("macro_type")]
        concentrations = []
        for r in range(183, 188):
            concentrations.append({
                "ranking": _cell(ws, r, "C"), "asset_type": _clean(_cell(ws, r, "D")), "position_type": _clean(_cell(ws, r, "E")),
                "market_code_type": _clean(_cell(ws, r, "F")), "market_code": _clean(_cell(ws, r, "G")),
                "aggregated_value": _num(_cell(ws, r, "H")), "aggregated_value_pct": _num(_cell(ws, r, "I")),
                "counterparty_name": _clean(_cell(ws, r, "J")), "counterparty_lei": _clean(_cell(ws, r, "K")), "counterparty_bic": _clean(_cell(ws, r, "L")),
            })
        aif["portfolio_concentrations"] = [c for c in concentrations if c.get("asset_type")]
        markets3 = []
        for r in range(194, 197):
            markets3.append({"ranking": _cell(ws, r, "C"), "market_code_type": _clean(_cell(ws, r, "E")), "market_code": _clean(_cell(ws, r, "G")), "aggregated_value": _num(_cell(ws, r, "I"))})
        aif["aif_principal_markets"] = [m for m in markets3 if m.get("market_code_type")]
        # 5 main instruments (rows 141-145).
        instruments5 = []
        for r in range(141, 146):
            instruments5.append({
                "ranking": _cell(ws, r, "C"), "sub_asset_type": _clean(_cell(ws, r, "D")), "instrument_code_type": _clean(_cell(ws, r, "E")),
                "instrument_name": _clean(_cell(ws, r, "F")), "isin": _clean(_cell(ws, r, "G")),
                "aii_exchange_code": _clean(_cell(ws, r, "H")), "position_type": _clean(_cell(ws, r, "I")),
                "position_value": _num(_cell(ws, r, "J")), "hedging_rate": _num(_cell(ws, r, "L")),
            })
        aif["main_instruments"] = [i for i in instruments5 if i.get("sub_asset_type")]

    if TEMPLATE_SHEET_AIF2 in wb.sheetnames:
        ws = wb[TEMPLATE_SHEET_AIF2]
        aif.update(simple(ws, {
            137: "annual_investment_return_rate",
            148: "trading_pct_securities_exchange", 149: "trading_pct_securities_otc",
            150: "trading_pct_derivatives_exchange", 151: "trading_pct_derivatives_otc",
            152: "derivatives_ccp_rate", 153: "derivatives_bilateral_rate",
            154: "repo_ccp_rate", 155: "repo_bilateral_rate", 156: "repo_triparty_rate",
            157: "collateral_cash_posted", 158: "collateral_securities_posted", 159: "collateral_other_posted",
            172: "ccp_clearing_flag", 185: "unencumbered_cash", 193: "provide_withdrawal_rights",
            194: "redemption_frequency", 195: "redemption_notice_period_days", 196: "redemption_lockup_period_days",
            197: "side_pocket_rate", 198: "gates_rate", 199: "dealing_suspension_rate",
            200: "other_arrangement_type", 201: "other_arrangement_rate", 202: "total_arrangement_rate",
            203: "investor_preferential_treatment_flag", 204: "pref_disclosure_flag", 205: "pref_liquidity_flag",
            206: "pref_fee_flag", 207: "pref_other_flag", 210: "total_financing_amount", 218: "total_open_positions",
            279: "stress_test_article15", 280: "stress_test_article16",
        }, max_row=379))
        portfolio_liq = {code: _num(_cell(ws, r, "C")) for r, code in zip(range(219, 226), [c for c, _ in LIQUIDITY_BUCKETS])}
        aif["portfolio_liquidity"] = portfolio_liq
        investor_liq = {code: _num(_cell(ws, r, "C")) for r, code in zip(range(229, 236), [c for c, _ in LIQUIDITY_BUCKETS])}
        aif["investor_liquidity"] = investor_liq
        financing = {code: _num(_cell(ws, r, "C")) for r, code in zip(range(275, 282), [c for c, _ in LIQUIDITY_BUCKETS])}
        aif["financing_buckets"] = financing
        investor_groups = []
        for r, (code, _label) in zip(range(262, 272), INVESTOR_GROUP_TYPES.items()):
            investor_groups.append({"code": code, "rate": _num(_cell(ws, r, "E"))})
        aif["investor_groups"] = investor_groups
        historical = []
        for i, m in enumerate(MONTHS):
            historical.append({
                "month": m,
                "gross_return": _num(_cell(ws, 286 + i, "C")),
                "net_return": _num(_cell(ws, 299 + i, "C")),
                "nav_change": _num(_cell(ws, 312 + i, "C")),
                "subscriptions": _num(_cell(ws, 325 + i, "C")),
                "redemptions": _num(_cell(ws, 338 + i, "C")),
            })
        aif["historical_performance"] = historical
        risk_measures = []
        for r, code in zip(range(160, 167), ["NET_EQTY_DELTA", "NET_DV01", "NET_CS01", "VEGA_EXPO", "VAR", "NET_FX_DELTA", "NET_CTY_DELTA"]):
            risk_measures.append({"code": code, "value": _num(_cell(ws, r, "D"))})
        aif["market_risk_measures"] = risk_measures
        instrument_exposures = []
        for r in range(12, 104):
            code = _cell(ws, r, "C")
            if code:
                instrument_exposures.append({"code": code, "gross_value": _num(_cell(ws, r, "E")), "long_value": _num(_cell(ws, r, "G")), "short_value": _num(_cell(ws, r, "H"))})
        aif["instrument_exposures"] = instrument_exposures
        asset_turnovers = []
        for r in range(108, 142):
            code = _cell(ws, r, "C")
            if code:
                asset_turnovers.append({"code": code, "market_value": _num(_cell(ws, r, "E")), "notional_value": _num(_cell(ws, r, "G"))})
        aif["asset_turnovers"] = asset_turnovers
        counterparties_to = []
        for r in range(194, 199):
            counterparties_to.append({"ranking": _cell(ws, r, "C"), "flag": _clean(_cell(ws, r, "D")), "name": _clean(_cell(ws, r, "E")), "lei": _clean(_cell(ws, r, "F")), "bic": _clean(_cell(ws, r, "G")), "nav_pct": _num(_cell(ws, r, "H"))})
        aif["counterparty_exposures_to_aif"] = [c for c in counterparties_to if c.get("name")]
        counterparties_from = []
        for r in range(202, 207):
            counterparties_from.append({"ranking": _cell(ws, r, "C"), "flag": _clean(_cell(ws, r, "D")), "name": _clean(_cell(ws, r, "E")), "lei": _clean(_cell(ws, r, "F")), "bic": _clean(_cell(ws, r, "G")), "nav_pct": _num(_cell(ws, r, "H"))})
        aif["counterparty_exposures_from_aif"] = [c for c in counterparties_from if c.get("name")]
        ccps = []
        for r in range(213, 216):
            ccps.append({"ranking": _cell(ws, r, "C"), "name": _clean(_cell(ws, r, "D")), "lei": _clean(_cell(ws, r, "E")), "bic": _clean(_cell(ws, r, "F")), "exposure_value": _num(_cell(ws, r, "G"))})
        aif["ccp_exposures"] = [c for c in ccps if c.get("name")]

    if TEMPLATE_SHEET_AIF4 in wb.sheetnames:
        ws = wb[TEMPLATE_SHEET_AIF4]
        aif.update(simple(ws, {
            6: "rehypothecation_flag", 7: "rehypothecation_rate", 10: "unsecured_borrowing_amount",
            11: "secured_borrowing_prime_brokerage_amount", 12: "secured_borrowing_reverse_repo_amount",
            13: "secured_borrowing_other_amount", 15: "exchange_traded_derivatives_exposure", 16: "otc_derivatives_amount",
            17: "short_position_borrowed_securities_value", 26: "gross_method_rate", 27: "commitment_method_rate",
        }, max_row=37))
        controlled = []
        for r in range(21, 24):
            controlled.append({"name": _clean(_cell(ws, r, "C")), "lei": _clean(_cell(ws, r, "D")), "bic": _clean(_cell(ws, r, "F")), "exposure_value": _num(_cell(ws, r, "H"))})
        aif["controlled_structures"] = [c for c in controlled if c.get("name")]
        sources = []
        for r in range(31, 36):
            sources.append({"ranking": _cell(ws, r, "C"), "flag": _clean(_cell(ws, r, "D")), "name": _clean(_cell(ws, r, "E")), "lei": _clean(_cell(ws, r, "F")), "bic": _clean(_cell(ws, r, "G")), "leverage_amount": _num(_cell(ws, r, "H"))})
        aif["borrowing_sources"] = [s for s in sources if s.get("name")]

    if not aifm and not aif:
        errors.append("No recognisable data found — is this the official ESMA consolidated template?")

    return aifm, aif, errors


def _as_bool(value: Any) -> bool | None:
    if value is None or value in ("", "—", "-"):
        return None
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "yes", "y")


def default_market_risk_measures() -> list[dict]:
    return [{"code": code, "value": None} for code, _label, _mandatory in MARKET_RISK_MEASURE_TYPES]


def default_instrument_exposures() -> list[dict]:
    return [{"code": code, "gross_value": None, "long_value": None, "short_value": None} for code, _label in SUB_ASSET_TYPES]


def default_asset_turnovers() -> list[dict]:
    return [{"code": code, "market_value": None, "notional_value": None} for code, _label in TURNOVER_SUB_ASSET_TYPES]


def default_investor_groups() -> list[dict]:
    return [{"code": code, "rate": None} for code in INVESTOR_GROUP_TYPES]


def default_historical_performance() -> list[dict]:
    return [{"month": m, "gross_return": None, "net_return": None, "nav_change": None, "subscriptions": None, "redemptions": None} for m in MONTHS]


def submission_from_dicts(
    aifm_info: dict[str, Any], aif_info: dict[str, Any], period_start: date, period_end: date,
) -> AifmdSubmission:
    """Builds an AifmdSubmission from the loosely-typed dicts produced by
    parse_aifmd_workbook() (Auto mode) — shared with the Manual-mode form's
    own draft-building helper for a single source of truth on field names."""

    def f(d: dict, key: str, default: Any = None) -> Any:
        v = d.get(key, default)
        return default if v in (None, "") else v

    sub = AifmdSubmission(
        sender_entity_type=str(f(aifm_info, "sender_entity_type", "S")),
        sender_id=str(f(aifm_info, "sender_id", "")),
        aifm_version=str(f(aifm_info, "aifm_version", SCHEMA_VERSION)),
        aifm_filing_type=str(f(aifm_info, "aifm_filing_type", "INIT")),
        aifm_content_type=str(f(aifm_info, "aifm_content_type", "1")),
        reporting_period_type=str(f(aifm_info, "reporting_period_type", "Y1")),
        reporting_period_year=int(f(aifm_info, "reporting_period_year", period_end.year)),
        reporting_period_start=period_start,
        reporting_period_end=period_end,
        aifm_change_frequency_code=f(aifm_info, "aifm_change_frequency_code"),
        aifm_change_contents_code=f(aifm_info, "aifm_change_contents_code"),
        aifm_change_quarter=f(aifm_info, "aifm_change_quarter"),
        aifm_last_reporting_flag=bool(_as_bool(aifm_info.get("aifm_last_reporting_flag")) or False),
        aifm_assumptions=list(f(aifm_info, "aifm_assumptions", [])),
        aifm_reporting_code=int(f(aifm_info, "aifm_reporting_code", 1)),
        aifm_jurisdiction=str(f(aifm_info, "aifm_jurisdiction", "LU")),
        aifm_national_code=str(f(aifm_info, "aifm_national_code", "")),
        aifm_name=str(f(aifm_info, "aifm_name", "")),
        aifm_eea_flag=bool(_as_bool(aifm_info.get("aifm_eea_flag")) if aifm_info.get("aifm_eea_flag") is not None else True),
        aifm_no_reporting_flag=bool(_as_bool(aifm_info.get("aifm_no_reporting_flag")) or False),
        aifm_lei=str(f(aifm_info, "aifm_lei", "")),
        aifm_bic=str(f(aifm_info, "aifm_bic", "")),
        aifm_old_identifier_country=str(f(aifm_info, "aifm_old_identifier_country", "")),
        aifm_old_identifier_code=str(f(aifm_info, "aifm_old_identifier_code", "")),
        aifm_principal_markets=list(f(aifm_info, "aifm_principal_markets", [])),
        aifm_principal_instruments=list(f(aifm_info, "aifm_principal_instruments", [])),
        total_aum_eur=float(f(aifm_info, "total_aum_eur", 0.0)),
        total_aum_base_currency=float(f(aifm_info, "total_aum_base_currency", 0.0)),
        aifm_base_currency=str(f(aifm_info, "aifm_base_currency", "EUR")).upper(),
        aifm_fx_ref_rate_type=str(f(aifm_info, "aifm_fx_ref_rate_type", "")),
        eur_fx_rate=float(f(aifm_info, "eur_fx_rate", 1.0)),
        aifm_fx_ref_rate_description=str(f(aifm_info, "aifm_fx_ref_rate_description", "")),
        aif_version=str(f(aif_info, "aif_version", SCHEMA_VERSION)),
        aif_filing_type=str(f(aif_info, "aif_filing_type", "INIT")),
        aif_content_type=str(f(aif_info, "aif_content_type", "2")),
        aif_change_frequency_code=f(aif_info, "aif_change_frequency_code"),
        aif_change_contents_code=f(aif_info, "aif_change_contents_code"),
        aif_change_quarter=f(aif_info, "aif_change_quarter"),
        aif_last_reporting_flag=bool(_as_bool(aif_info.get("aif_last_reporting_flag")) or False),
        aif_assumptions=list(f(aif_info, "aif_assumptions", [])),
        aif_legal_form=str(f(aif_info, "aif_legal_form", "O")),
        aif_national_number=str(f(aif_info, "aif_national_number", "")).zfill(8),
        aif_subfund_number=str(f(aif_info, "aif_subfund_number", "00000000")).zfill(8),
        aif_name=str(f(aif_info, "aif_name", "")),
        aif_eea_flag=bool(_as_bool(aif_info.get("aif_eea_flag")) if aif_info.get("aif_eea_flag") is not None else True),
        aif_reporting_code=int(f(aif_info, "aif_reporting_code", 1)),
        aif_domicile=str(f(aif_info, "aif_domicile", "")),
        aif_inception_date=aif_info.get("aif_inception_date") if isinstance(aif_info.get("aif_inception_date"), date) else None,
        aif_no_reporting_flag=bool(_as_bool(aif_info.get("aif_no_reporting_flag")) or False),
        aif_lei=str(f(aif_info, "aif_lei", "")),
        aif_isin=str(f(aif_info, "aif_isin", "")),
        aif_cusip=str(f(aif_info, "aif_cusip", "")),
        aif_sedol=str(f(aif_info, "aif_sedol", "")),
        aif_bloomberg=str(f(aif_info, "aif_bloomberg", "")),
        aif_reuters=str(f(aif_info, "aif_reuters", "")),
        aif_ecb_code=str(f(aif_info, "aif_ecb_code", "")),
        aif_old_identifier_country=str(f(aif_info, "aif_old_identifier_country", "")),
        aif_old_identifier_code=str(f(aif_info, "aif_old_identifier_code", "")),
        aif_share_class_flag=bool(_as_bool(aif_info.get("aif_share_class_flag")) or False),
        aif_share_classes=list(f(aif_info, "aif_share_classes", [])),
        aif_master_feeder_status=str(f(aif_info, "aif_master_feeder_status", "NONE")),
        aif_master_aifs=list(f(aif_info, "aif_master_aifs", [])),
        aif_prime_brokers=list(f(aif_info, "aif_prime_brokers", [])),
        aif_aum_base_currency=float(f(aif_info, "aif_aum_base_currency", 0.0)),
        aif_base_currency=str(f(aif_info, "aif_base_currency", "EUR")).upper(),
        aif_fx_rate=float(f(aif_info, "aif_fx_rate", 1.0)),
        aif_fx_ref_rate_type=str(f(aif_info, "aif_fx_ref_rate_type", "")),
        aif_fx_ref_rate_description=str(f(aif_info, "aif_fx_ref_rate_description", "")),
        aif_nav=float(f(aif_info, "aif_nav", 0.0)),
        funding_source_country_1=str(f(aif_info, "funding_source_country_1", "")),
        funding_source_country_2=str(f(aif_info, "funding_source_country_2", "")),
        funding_source_country_3=str(f(aif_info, "funding_source_country_3", "")),
        predominant_aif_type=str(f(aif_info, "predominant_aif_type", "NONE")),
        investment_strategies=list(f(aif_info, "investment_strategies", [])),
        hft_transaction_number=f(aif_info, "hft_transaction_number"),
        hft_buy_sell_market_value=f(aif_info, "hft_buy_sell_market_value"),
        main_instruments=list(f(aif_info, "main_instruments", [])),
        nav_geo_focus=dict(f(aif_info, "nav_geo_focus", {})),
        aum_geo_focus=dict(f(aif_info, "aum_geo_focus", {})),
        ten_principal_exposures=list(f(aif_info, "ten_principal_exposures", [])),
        portfolio_concentrations=list(f(aif_info, "portfolio_concentrations", [])),
        typical_position_size=str(f(aif_info, "typical_position_size", "")),
        aif_principal_markets=list(f(aif_info, "aif_principal_markets", [])),
        investor_beneficial_owners_rate=f(aif_info, "investor_beneficial_owners_rate"),
        investor_professional_rate=f(aif_info, "investor_professional_rate"),
        investor_retail_rate=f(aif_info, "investor_retail_rate"),
        instrument_exposures=list(f(aif_info, "instrument_exposures", [])),
        asset_turnovers=list(f(aif_info, "asset_turnovers", [])),
        currency_exposures=list(f(aif_info, "currency_exposures", [])),
        dominant_influence_companies=list(f(aif_info, "dominant_influence_companies", [])),
        annual_investment_return_rate=str(f(aif_info, "annual_investment_return_rate", "")),
        market_risk_measures=list(f(aif_info, "market_risk_measures", [])),
        trading_pct_securities_exchange=f(aif_info, "trading_pct_securities_exchange"),
        trading_pct_securities_otc=f(aif_info, "trading_pct_securities_otc"),
        trading_pct_derivatives_exchange=f(aif_info, "trading_pct_derivatives_exchange"),
        trading_pct_derivatives_otc=f(aif_info, "trading_pct_derivatives_otc"),
        derivatives_ccp_rate=f(aif_info, "derivatives_ccp_rate"),
        derivatives_bilateral_rate=f(aif_info, "derivatives_bilateral_rate"),
        repo_ccp_rate=f(aif_info, "repo_ccp_rate"),
        repo_bilateral_rate=f(aif_info, "repo_bilateral_rate"),
        repo_triparty_rate=f(aif_info, "repo_triparty_rate"),
        collateral_cash_posted=f(aif_info, "collateral_cash_posted"),
        collateral_securities_posted=f(aif_info, "collateral_securities_posted"),
        collateral_other_posted=f(aif_info, "collateral_other_posted"),
        counterparty_exposures_to_aif=list(f(aif_info, "counterparty_exposures_to_aif", [])),
        counterparty_exposures_from_aif=list(f(aif_info, "counterparty_exposures_from_aif", [])),
        ccp_clearing_flag=_as_bool(aif_info.get("ccp_clearing_flag")),
        ccp_exposures=list(f(aif_info, "ccp_exposures", [])),
        portfolio_liquidity=dict(f(aif_info, "portfolio_liquidity", {})),
        unencumbered_cash=f(aif_info, "unencumbered_cash"),
        investor_liquidity=dict(f(aif_info, "investor_liquidity", {})),
        provide_withdrawal_rights=_as_bool(aif_info.get("provide_withdrawal_rights")),
        redemption_frequency=str(f(aif_info, "redemption_frequency", "")),
        redemption_notice_period_days=f(aif_info, "redemption_notice_period_days"),
        redemption_lockup_period_days=f(aif_info, "redemption_lockup_period_days"),
        side_pocket_rate=f(aif_info, "side_pocket_rate"),
        gates_rate=f(aif_info, "gates_rate"),
        dealing_suspension_rate=f(aif_info, "dealing_suspension_rate"),
        other_arrangement_type=str(f(aif_info, "other_arrangement_type", "")),
        other_arrangement_rate=f(aif_info, "other_arrangement_rate"),
        total_arrangement_rate=f(aif_info, "total_arrangement_rate"),
        investor_preferential_treatment_flag=_as_bool(aif_info.get("investor_preferential_treatment_flag")),
        pref_disclosure_flag=_as_bool(aif_info.get("pref_disclosure_flag")),
        pref_liquidity_flag=_as_bool(aif_info.get("pref_liquidity_flag")),
        pref_fee_flag=_as_bool(aif_info.get("pref_fee_flag")),
        pref_other_flag=_as_bool(aif_info.get("pref_other_flag")),
        investor_groups=list(f(aif_info, "investor_groups", [])),
        total_financing_amount=f(aif_info, "total_financing_amount"),
        financing_buckets=dict(f(aif_info, "financing_buckets", {})),
        total_open_positions=f(aif_info, "total_open_positions"),
        historical_performance=list(f(aif_info, "historical_performance", [])),
        stress_test_article15=str(f(aif_info, "stress_test_article15", "")),
        stress_test_article16=str(f(aif_info, "stress_test_article16", "")),
        rehypothecation_flag=_as_bool(aif_info.get("rehypothecation_flag")),
        rehypothecation_rate=f(aif_info, "rehypothecation_rate"),
        unsecured_borrowing_amount=f(aif_info, "unsecured_borrowing_amount"),
        secured_borrowing_prime_brokerage_amount=f(aif_info, "secured_borrowing_prime_brokerage_amount"),
        secured_borrowing_reverse_repo_amount=f(aif_info, "secured_borrowing_reverse_repo_amount"),
        secured_borrowing_other_amount=f(aif_info, "secured_borrowing_other_amount"),
        exchange_traded_derivatives_exposure=f(aif_info, "exchange_traded_derivatives_exposure"),
        otc_derivatives_amount=f(aif_info, "otc_derivatives_amount"),
        short_position_borrowed_securities_value=f(aif_info, "short_position_borrowed_securities_value"),
        controlled_structures=list(f(aif_info, "controlled_structures", [])),
        gross_method_rate=f(aif_info, "gross_method_rate"),
        commitment_method_rate=f(aif_info, "commitment_method_rate"),
        borrowing_sources=list(f(aif_info, "borrowing_sources", [])),
    )

    # Recover the legal-form prefix + national number from the raw item-17
    # code, when the workbook (rather than the manual form) is the source.
    raw_code = aif_info.get("aif_national_code_full")
    if raw_code and not sub.aif_national_number:
        m = re.fullmatch(r"([OKV])(\d{8})", str(raw_code).strip())
        if m:
            sub.aif_legal_form, sub.aif_national_number = m.group(1), m.group(2)

    if not sub.market_risk_measures:
        sub.market_risk_measures = default_market_risk_measures()
    if not sub.instrument_exposures:
        sub.instrument_exposures = default_instrument_exposures()
    if not sub.asset_turnovers:
        sub.asset_turnovers = default_asset_turnovers()
    if not sub.investor_groups:
        sub.investor_groups = default_investor_groups()
    if not sub.historical_performance:
        sub.historical_performance = default_historical_performance()

    return sub
