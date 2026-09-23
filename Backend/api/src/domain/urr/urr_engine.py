"""URR (CSSF U 1.1) reporting engine.

U 1.1 is the CSSF's monthly financial/statistical report on Luxembourg
UCIs (UCITS, Part II UCIs, SIFs, SICARs — Circular CSSF 15/627), filed once
per sub-fund (no consolidated report for the umbrella entity), 10 calendar
days after each month end, as XML via eDesk.

The field list, codes and exact XML element names below are taken directly
from a real CSSF example workbook (U1.1/Examples_U11_2025.xlsx, version 5.1
- 12/2025 of the U1.1 reporting spec) — not guessed. The enumerated code
lists below (NAV/redemption frequency, open/closed-end, investor base,
reason for closing) are taken directly from the currently-effective CSSF
schema (U11.zip -> cssf.datatypes.xsd, "applicable as of the 12/2025
reference period") — the enumeration values and their documentation text
are copied verbatim, including the source's own "semy-weekly" / "UCIS"
spelling, rather than silently corrected.
"""

from __future__ import annotations

import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date
from typing import Any
from xml.dom import minidom

U11_NAMESPACE = "http://www.cssf.lu/2015/U11"
U11_SCHEMA_LOCATION = "http://www.cssf.lu/2015/U11 U11.xsd"

REPORTING_STATUSES = ["final report", "null report"]

# Code -> label, from cssf.datatypes.xsd's "NavFrequency" xs:simpleType
# (positiveInteger enumeration, values 1-20). Source spelling "semy-weekly"
# (values 3/4) preserved verbatim.
NAV_FREQUENCY_CODES: dict[int, str] = {
    1: "Daily", 2: "Daily and month-end", 3: "Semy-weekly", 4: "Semy-weekly and month-end",
    5: "Weekly", 6: "Weekly and month-end", 7: "Weekly and first day of the month",
    8: "Weekly and for each subscription or redemption", 9: "Semi-monthly",
    10: "Semi-monthly and month end", 11: "Semi-monthly and quarter end",
    12: "Semi-monthly and for each subscription or redemption", 13: "Monthly",
    14: "Monthly and for each subscription or redemption", 15: "Quarterly",
    16: "Quarterly and for each subscription or redemption", 17: "Semi-annual",
    18: "Semi-annual and for each subscription or redemption", 19: "Annual",
    20: "Annual and for each subscription or redemption",
}

# "NavRedemptionFreq" xs:simpleType (positiveInteger enumeration, 1-9).
NAV_REDEMPTION_FREQ_CODES: dict[int, str] = {
    1: "Daily", 2: "Weekly", 3: "Semi-monthly", 4: "Monthly", 5: "Quarterly",
    6: "Biannual", 7: "Annual", 8: "Only at the end of maturity", 9: "Other/Irregular",
}

# "OpenClosedEnd" xs:simpleType (positiveInteger enumeration, 1-2).
OPEN_CLOSED_END_CODES: dict[int, str] = {
    1: "Open-end investment funds", 2: "Closed-end investment funds",
}

# "InvestorBase" xs:simpleType (positiveInteger enumeration, 1-4).
INVESTOR_BASE_CODES: dict[int, str] = {
    1: "Professional investor", 2: "Well-informed investor", 3: "Retail investor", 4: "Mixed",
}

# "CssfReasonForClosingUCI" xs:simpleType (positiveInteger enumeration,
# 1-5) — UCI level (field 2090). Source documentation text "The UCIS is"
# (values 4/5) preserved verbatim.
UCI_CLOSING_REASON_CODES: dict[int, str] = {
    1: "Put into liquidation",
    2: "Merger into another UCI (wherever it is domiciled)",
    3: "The UCI status is given up and the UCI has adopted a legal status different from UCI",
    4: "The UCIS is closed temporarily (see CSSF circular 12/540)",
    5: "The UCIS is closed definitively (see CSSF circular 12/540)",
}

# "CssfReasonForClosingSC" xs:simpleType (positiveInteger enumeration,
# 1-4) — unit/share class level (field 4070).
SHARE_CLASS_CLOSING_REASON_CODES: dict[int, str] = {
    1: "Put into liquidation",
    2: "Merger into another unit / share class (wherever it is domiciled)",
    3: "The UCI status is given up and the UCI has adopted a legal status different from UCI",
    4: "The unit / share class is closed temporarily",
}


# (key, label) — labels copied verbatim, including the field code, from
# Annex 2 ("U1.1 reporting item description") of the CSSF's Guidelines on
# the U1.1 reporting.
SHARE_CLASS_FIELDS = [
    ("cssf_code", "4010 — CSSF code of the unit/share class"),
    ("name", "4020 — Name of the unit/share class"),
    ("currency", "4030 — Base currency of the unit/share class"),
    ("isin", "4040 — ISIN code"),
    ("launch_date", "4050 — Launch date of the unit/share class"),
    ("closing_date", "4060 — Closing date of the unit/share class"),
    ("closing_reason", "4070 — Reason for closing of the unit/share class"),
    ("spec_fx", "4080 — Unit/share class with specific instruments – FX"),
    ("spec_ir", "4090 — Unit/share class with specific instruments – Interest Rate"),
    ("spec_other", "4100 — Unit/share class with specific instruments – Other"),
    ("investor_base", "4110 — Investor base"),
    ("units_outstanding", "5010 — Number of outstanding units/shares"),
    ("nav_per_unit_uci_ccy", "5020 — Net asset value per unit/share in the base currency of the UCI"),
    ("nav_per_unit_class_ccy", "5030 — Net asset value per unit/share in the base currency of the unit/share class"),
    ("net_return_pct", "5040 — Net return per unit/share in the base currency of the unit/share class"),
    ("units_issued", "5050 — Proceeds from units/shares issued in the base currency of the UCI"),
    ("units_redeemed", "5060 — Payments made for units/shares redeemed in the base currency of the UCI"),
    ("distributions", "5070 — Total distributions made in the base currency of the UCI"),
    ("distrib_per_unit_uci_ccy", "5080 — Amount of distributions made per unit/share in the base currency of the UCI"),
    ("distrib_per_unit_class_ccy", "5090 — Amount of distributions made per unit/share in the base currency of the unit/share class"),
    ("income_uci_ccy", "5100 — Income in the base currency of the UCI"),
    ("expenses_uci_ccy", "5110 — Expenses in the base currency of the UCI"),
]

SHARE_CLASS_DEFAULT_ROW: dict[str, Any] = {
    "cssf_code": "", "name": "", "currency": "EUR", "isin": "", "launch_date": None,
    "closing_date": None, "closing_reason": None, "spec_fx": False, "spec_ir": False, "spec_other": False,
    "investor_base": 1, "units_outstanding": 0.0, "nav_per_unit_uci_ccy": 0.0, "nav_per_unit_class_ccy": 0.0,
    "net_return_pct": 0.0, "units_issued": 0.0, "units_redeemed": 0.0, "distributions": 0.0,
    "distrib_per_unit_uci_ccy": 0.0, "distrib_per_unit_class_ccy": 0.0, "income_uci_ccy": 0.0, "expenses_uci_ccy": 0.0,
}


@dataclass
class UrrSubmission:
    # Section 1 — report/sender
    reference_month: date
    reporting_status: str
    sender_cssf_code: str
    employee_name: str
    employee_phone: str
    sender_email: str
    # Section 2 — UCI general info
    uci_cssf_code: str
    uci_name: str
    uci_currency: str
    lei: str = ""
    valuation_date: date | None = None
    nav_frequency: int = 1
    launch_date: date | None = None
    closing_date: date | None = None
    closing_reason: int | None = None
    redemption_frequency: int = 1
    open_closed_end: int = 1
    # Section 3 — UCI financial info
    nav_preceding_month: float = 0.0
    nav_this_month: float = 0.0
    proceeds_issued: float = 0.0
    payments_redeemed: float = 0.0
    total_distributions: float = 0.0
    # Sections 4+5 — per share class
    share_classes: list[dict[str, Any]] = field(default_factory=list)
    # Section 6 — investment income & expenses
    income: dict[str, float] = field(default_factory=dict)

    def validation_errors(self) -> list[str]:
        errors = []
        if not self.sender_cssf_code:
            errors.append("CSSF code of the sender (1030) is required.")
        if not self.employee_name:
            errors.append("Name of the employee (1040) is required.")
        if not self.sender_email:
            errors.append("Email address (1050) is required.")
        if not self.uci_cssf_code:
            errors.append("CSSF code of the UCI (2010) is required.")
        if not self.uci_name:
            errors.append("Name of the UCI (2020) is required.")
        if len(self.uci_currency) != 3:
            errors.append("Base currency of the UCI (2030) must be a 3-letter ISO code.")
        if self.reporting_status == "final report":
            if not self.share_classes:
                errors.append("At least one share class (section 4/5) is required for a final report.")
            for i, sc in enumerate(self.share_classes, start=1):
                if not sc.get("cssf_code"):
                    errors.append(f"Share class #{i}: CSSF code (4010) is required.")
                if not sc.get("isin"):
                    errors.append(f"Share class #{i}: ISIN (4040) is required.")
                elif not (len(sc["isin"]) == 12 and sc["isin"][:2].isalpha()):
                    errors.append(f"Share class #{i}: ISIN (4040) must be 2 letters + 10 characters.")
        if self.lei and len(self.lei) != 20:
            errors.append("LEI (2040) must be 20 characters when provided.")
        return errors


def _el(parent: ET.Element, tag: str, text: Any = None) -> ET.Element:
    e = ET.SubElement(parent, tag)
    if text is not None and text != "":
        e.text = str(text)
    return e


def _fmt_date(d: date | None) -> str:
    return d.isoformat() if d else ""


def _fmt_amount(v: float, decimals: int = 5) -> str:
    return f"{float(v):.{decimals}f}"


def build_u11_xml(sub: UrrSubmission) -> str:
    """Builds the U 1.1 XML exactly matching the structure/tag names of the
    real CSSF example (final report + null report variants)."""
    root = ET.Element("U11", {
        "xmlns": U11_NAMESPACE,
        "xmlns:xsi": "http://www.w3.org/2001/XMLSchema-instance",
        "version": "1.0",
        "creationDateTime": date.today().isoformat() + "T00:00:00Z",
        "xsi:schemaLocation": U11_SCHEMA_LOCATION,
    })

    gen = _el(root, "GeneralInformation")
    _el(gen, "ReferenceMonth", sub.reference_month.strftime("%Y%m"))
    _el(gen, "ReportingStatus", sub.reporting_status)
    _el(gen, "SenderCssfCode", sub.sender_cssf_code)
    _el(gen, "TypeOfTheSender")
    _el(gen, "CodeOfTheSender")
    senders = _el(gen, "SenderInformations")
    for name_phone in sub.employee_name.split(";"):
        name_phone = name_phone.strip()
        if not name_phone:
            continue
        si = _el(senders, "SenderInformation")
        _el(si, "NameOfEmployee", name_phone.split("-")[0].strip() if "-" in name_phone else name_phone)
        _el(si, "PhoneNumber", sub.employee_phone)
        _el(si, "EmailAdress", sub.sender_email)

    uci_gen = _el(root, "UCIGeneralInformation")
    basic = _el(uci_gen, "UCISignaBasic")
    _el(basic, "UCICssfCode", sub.uci_cssf_code)
    _el(basic, "TypeOfTheDeclarant")
    _el(basic, "CodeOfTheDeclarant")
    _el(basic, "UCIName", sub.uci_name)
    _el(basic, "UCIBaseCurrency", sub.uci_currency)
    if sub.lei:
        _el(basic, "LEI", sub.lei)

    if sub.reporting_status == "null report":
        return _pretty(root)

    sdate = _el(uci_gen, "UCISignaDate")
    _el(sdate, "ValuationDate", _fmt_date(sub.valuation_date))
    nav_info = _el(sdate, "NavInformation")
    _el(nav_info, "NavFreq", sub.nav_frequency)
    _el(nav_info, "NavRedemptionFreq", sub.redemption_frequency)
    _el(nav_info, "OpenClosedEnd", sub.open_closed_end)
    if sub.launch_date:
        _el(sdate, "LaunchDate", _fmt_date(sub.launch_date))
    if sub.closing_date:
        _el(sdate, "ClosingDate", _fmt_date(sub.closing_date))
    if sub.closing_reason:
        _el(sdate, "ReasonForClosing", sub.closing_reason)

    fin = _el(root, "UCIFinancialInformation")
    _el(fin, "NAVPrecedingRefMonth", _fmt_amount(sub.nav_preceding_month))
    _el(fin, "NAVRefMonth", _fmt_amount(sub.nav_this_month))
    _el(fin, "TotUnitsOrSharesIssued", _fmt_amount(sub.proceeds_issued))
    _el(fin, "TotUnitsOrSharesRedeemed", _fmt_amount(sub.payments_redeemed))
    _el(fin, "TotalDistributions", _fmt_amount(sub.total_distributions))

    classes_el = _el(root, "Unit_Share_Informations")
    for sc in sub.share_classes:
        usi = _el(classes_el, "Unit_Share_Information")
        _el(usi, "CssfCode", sc.get("cssf_code", ""))
        g = _el(usi, "UnitShareClassGeneralInformation")
        _el(g, "Unit_ShareClass_Name", sc.get("name", ""))
        _el(g, "BaseCurrency", sc.get("currency", ""))
        _el(g, "ISINCode", sc.get("isin", ""))
        if sc.get("launch_date"):
            _el(g, "LaunchDate", _fmt_date(sc["launch_date"]))
        if sc.get("closing_date"):
            _el(g, "ClosingDate", _fmt_date(sc["closing_date"]))
        if sc.get("closing_reason"):
            _el(g, "ReasonForClosing", sc["closing_reason"])
        _el(g, "Unit_ShareClass_SpecInstr_FX", str(bool(sc.get("spec_fx"))).lower())
        _el(g, "Unit_ShareClass_SpecInstr_InterestRate", str(bool(sc.get("spec_ir"))).lower())
        _el(g, "Unit_ShareClass_SpecInstr_Other", str(bool(sc.get("spec_other"))).lower())
        _el(g, "InvestorBase", sc.get("investor_base", 1))

        f = _el(usi, "UnitShareClassFinancialInformation")
        _el(f, "NumberUnitOrShareOutstanding", _fmt_amount(sc.get("units_outstanding", 0)))
        _el(f, "NetAssetValueInSubFundCurrency", _fmt_amount(sc.get("nav_per_unit_uci_ccy", 0), 10))
        _el(f, "NetAssetValueInShareClassCurrency", _fmt_amount(sc.get("nav_per_unit_class_ccy", 0), 10))
        _el(f, "NetReturn", _fmt_amount(sc.get("net_return_pct", 0)))
        _el(f, "UnitOrSharesIssued", _fmt_amount(sc.get("units_issued", 0)))
        _el(f, "UnitOrSharesRedeemed", _fmt_amount(sc.get("units_redeemed", 0)))
        _el(f, "Distributions", _fmt_amount(sc.get("distributions", 0)))
        _el(f, "AmountDistributionsInUCICurrency", _fmt_amount(sc.get("distrib_per_unit_uci_ccy", 0)))
        _el(f, "AmountDistributionsInShareClassCurrency", _fmt_amount(sc.get("distrib_per_unit_class_ccy", 0)))
        _el(f, "IncomeInFundCurrency", _fmt_amount(sc.get("income_uci_ccy", 0)))
        _el(f, "ExpensesInFundCurrency", _fmt_amount(sc.get("expenses_uci_ccy", 0)))

    inc = sub.income
    ie = _el(root, "UCI_Inv_Income_and_Expenses")
    d1 = _el(ie, "DividendsOnShares")
    _el(d1, "DividendsOnShares", _fmt_amount(inc.get("dividends_shares", 0)))
    _el(d1, "DivShOfWhichFromLu", _fmt_amount(inc.get("dividends_shares_lu", 0)))
    d2 = _el(ie, "DividendsOnUCI_Units_Shares")
    _el(d2, "DividendsOnUCI_Units_Shares", _fmt_amount(inc.get("dividends_uci_units", 0)))
    _el(d2, "DivUCIOfWhichFromLu", _fmt_amount(inc.get("dividends_uci_units_lu", 0)))
    d3 = _el(ie, "Interest")
    _el(d3, "Interest_bonds_debt_securities", _fmt_amount(inc.get("interest_bonds", 0)))
    _el(d3, "DivIntBondsOfWhichFromLu", _fmt_amount(inc.get("interest_bonds_lu", 0)))
    _el(ie, "InterestLoansBanksBearingAccounts", _fmt_amount(inc.get("interest_loans_banks", 0)))
    d4 = _el(ie, "BankInterest")
    _el(d4, "BankInterest", _fmt_amount(inc.get("bank_interest", 0)))
    _el(d4, "BankIntOfWhichFromLu", _fmt_amount(inc.get("bank_interest_lu", 0)))
    d5 = _el(ie, "OtherInterest")
    _el(d5, "OtherInterest", _fmt_amount(inc.get("other_interest", 0)))
    _el(d5, "OtherIntOfWhichFromLu", _fmt_amount(inc.get("other_interest_lu", 0)))
    _el(ie, "OtherIncome", _fmt_amount(inc.get("other_income", 0)))
    d6 = _el(ie, "OtherIncomeDetails")
    _el(d6, "RentalIncome", _fmt_amount(inc.get("rental_income", 0)))
    _el(d6, "CommissionIncome", _fmt_amount(inc.get("commission_income", 0)))
    _el(d6, "OtherOfWhichFromOtherIncome", _fmt_amount(inc.get("other_income_other", 0)))
    _el(ie, "Charges", _fmt_amount(inc.get("charges", 0)))
    d7 = _el(ie, "ChargesDetails")
    _el(d7, "AdvisoryManagementFees", _fmt_amount(inc.get("advisory_mgmt_fees", 0)))
    _el(d7, "DepositaryBankFees", _fmt_amount(inc.get("depositary_fees", 0)))
    _el(d7, "AdministrationFees", _fmt_amount(inc.get("admin_fees", 0)))
    _el(d7, "AuditInspectionFees", _fmt_amount(inc.get("audit_fees", 0)))
    _el(d7, "DistributionFees", _fmt_amount(inc.get("distribution_fees", 0)))
    _el(d7, "OtherAdministrativeExpenses", _fmt_amount(inc.get("other_admin_expenses", 0)))
    _el(d7, "SubscriptionTax", _fmt_amount(inc.get("subscription_tax", 0)))
    _el(d7, "OtherTaxes", _fmt_amount(inc.get("other_taxes", 0)))
    _el(ie, "PerformanceFees", _fmt_amount(inc.get("performance_fees", 0)))
    _el(ie, "OtherExpenses", _fmt_amount(inc.get("other_expenses", 0)))
    d8 = _el(ie, "OtherExpensesDetails")
    _el(d8, "InterestPaid", _fmt_amount(inc.get("interest_paid", 0)))
    _el(d8, "PersonnelExpenses", _fmt_amount(inc.get("personnel_expenses", 0)))
    _el(d8, "RentalExpenses", _fmt_amount(inc.get("rental_expenses", 0)))
    _el(d8, "Amortisation", _fmt_amount(inc.get("amortisation", 0)))
    _el(d8, "OtherOfWhichFromOtherExpenses", _fmt_amount(inc.get("other_expenses_other", 0)))
    _el(ie, "NetRealised", _fmt_amount(inc.get("net_realised", 0)))
    _el(ie, "NetUnrealised", _fmt_amount(inc.get("net_unrealised", 0)))

    return _pretty(root)


def _pretty(root: ET.Element) -> str:
    rough = ET.tostring(root, encoding="utf-8")
    return minidom.parseString(rough).toprettyxml(indent="  ", encoding="UTF-8").decode("utf-8")


# (key, label) — labels copied verbatim, including the field code, from
# Annex 2 ("U1.1 reporting item description") of the CSSF's Guidelines on
# the U1.1 reporting. Field 1040 ("Name of the employee and phone number")
# is one combined item in the regulation but two separate inputs here
# (name, phone) for usability — both cite the full official wording rather
# than inventing two separate field names for it.
FUND_INFO_FIELDS = [
    ("reference_month", "1010 — Reference month"),
    ("reporting_status", "1020 — Reporting status"),
    ("sender_cssf_code", "1030 — CSSF code of the sender"),
    ("employee_name", "1040 — Name of the employee and phone number (name)"),
    ("employee_phone", "1040 — Name of the employee and phone number (phone)"),
    ("sender_email", "1050 — Email address"),
    ("uci_cssf_code", "2010 — CSSF code of the UCI (CSSF code of the entity and CSSF code of the sub-fund)"),
    ("uci_name", "2020 — Name of the UCI (name of the entity and name of the sub-fund)"),
    ("uci_currency", "2030 — Base currency of the UCI"),
    ("lei", "2040 — LEI code"),
    ("valuation_date", "2050 — Valuation date"),
    ("nav_frequency", "2060 — NAV calculation frequency"),
    ("launch_date", "2070 — Launch date"),
    ("closing_date", "2080 — Closing date"),
    ("closing_reason", "2090 — Reason for closing"),
    ("redemption_frequency", "2100 — Redemption frequency"),
    ("open_closed_end", "2110 — Open- / Closed-end"),
    ("nav_preceding_month", "3010 — Total net asset value of the preceding reference month"),
    ("nav_this_month", "3020 — Total net asset value of the reference month"),
    ("proceeds_issued", "3030 — Proceeds from all units/shares issued during the reference month"),
    ("payments_redeemed", "3040 — Payments made for all units/shares redeemed during the reference month"),
    ("total_distributions", "3050 — Total distributions made during the reference month"),
]

# Short, single-line labels for the manual-entry grid (views.urr_ingestion)
# — the full Annex 2 wording above is always shown in that field's "?"
# tooltip, so no exactness is lost; this only keeps the 3-column grid from
# wrapping every label onto 2-3 lines (same pattern as the AIFMD module).
SHORT_LABELS: dict[str, str] = {
    "reference_month": "1010 — Reference month",
    "reporting_status": "1020 — Reporting status",
    "sender_cssf_code": "1030 — CSSF code of the sender",
    "employee_name": "1040 — Employee name",
    "employee_phone": "1040 — Employee phone",
    "sender_email": "1050 — Email address",
    "uci_cssf_code": "2010 — CSSF code of the UCI",
    "uci_name": "2020 — Name of the UCI",
    "uci_currency": "2030 — Base currency of the UCI",
    "lei": "2040 — LEI code",
    "valuation_date": "2050 — Valuation date",
    "nav_frequency": "2060 — NAV calculation frequency",
    "launch_date": "2070 — Launch date",
    "closing_date": "2080 — Closing date",
    "closing_reason": "2090 — Reason for closing",
    "redemption_frequency": "2100 — Redemption frequency",
    "open_closed_end": "2110 — Open-/Closed-end",
    "nav_preceding_month": "3010 — NAV (preceding month)",
    "nav_this_month": "3020 — NAV (reference month)",
    "proceeds_issued": "3030 — Proceeds from units issued",
    "payments_redeemed": "3040 — Payments for units redeemed",
    "total_distributions": "3050 — Total distributions",
}


def build_urr_template_workbook() -> bytes:
    """Ingestion template for URR/U 1.1:
    - 'Fund_Info': one row per section 1-3 / section 6 field (sections 1-3
      plus income/expenses) in long format — 'field' is the exact Annex 2
      code and label (e.g. '1010 — Reference month'), 'value' is what gets
      filled in. Long format (not one field per column) so the real,
      sometimes 80+ character Annex 2 wording can appear in full rather
      than being cut down to fit a column header.
    - 'Share_Classes': one row per class (sections 4-5), wide format since
      genuinely one row per class — column headers are the same exact
      Annex 2 code/label text.
    """
    import io

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    fill = PatternFill("solid", fgColor="0A2647")

    ws1 = wb.active
    ws1.title = "Fund_Info"
    header_cols = ["key", "field", "value"]
    for idx, col in enumerate(header_cols, start=1):
        cell = ws1.cell(row=1, column=idx, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 90
    ws1.column_dimensions["C"].width = 20
    for r, (key, label) in enumerate(FUND_INFO_FIELDS + INCOME_FIELDS, start=2):
        ws1.cell(row=r, column=1, value=key)
        ws1.cell(row=r, column=2, value=label)
        ws1.cell(row=r, column=3, value=None)
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("Share_Classes")
    for idx, (_key, label) in enumerate(SHARE_CLASS_FIELDS, start=1):
        cell = ws2.cell(row=1, column=idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        ws2.column_dimensions[cell.column_letter].width = max(18, min(40, len(label) + 2))
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_urr_workbook(file_bytes: bytes) -> tuple[dict[str, Any], list[dict[str, Any]], list[str]]:
    """Parses a workbook built from build_urr_template_workbook() back into
    (fund_info dict, share_classes list, errors)."""
    import io

    import pandas as pd

    errors: list[str] = []
    try:
        xls = pd.ExcelFile(io.BytesIO(file_bytes))
    except Exception as exc:
        return {}, [], [f"Could not read this workbook: {exc}"]

    fund_info: dict[str, Any] = {}
    if "Fund_Info" in xls.sheet_names:
        df = xls.parse("Fund_Info")
        for row in df.to_dict("records"):
            key = row.get("key")
            value = row.get("value")
            if key and pd.notna(value):
                fund_info[str(key)] = value
    else:
        errors.append("Sheet 'Fund_Info' not found.")

    share_classes: list[dict[str, Any]] = []
    if "Share_Classes" in xls.sheet_names:
        df = xls.parse("Share_Classes")
        label_to_key = {label: key for key, label in SHARE_CLASS_FIELDS}
        df = df.rename(columns=label_to_key)
        for row in df.to_dict("records"):
            if any(pd.notna(v) and str(v).strip() for v in row.values()):
                share_classes.append({k: (v if pd.notna(v) else "") for k, v in row.items()})
    else:
        errors.append("Sheet 'Share_Classes' not found.")

    return fund_info, share_classes, errors


# (key, label) — labels copied verbatim, including the field code, from
# Annex 2 of the CSSF's Guidelines on the U1.1 reporting.
INCOME_FIELDS = [
    ("dividends_shares", "6010 — Dividends on shares"),
    ("dividends_shares_lu", "6011 — of which from Luxembourg"),
    ("dividends_uci_units", "6020 — Dividends on UCI units/shares"),
    ("dividends_uci_units_lu", "6021 — of which from Luxembourg"),
    ("interest_bonds", "6030 — Interest on bonds and other debt securities"),
    ("interest_bonds_lu", "6031 — of which from Luxembourg"),
    ("interest_loans_banks", "6040 — Interest on loans, bank and other interest-bearing accounts"),
    ("bank_interest", "6041 — Bank interest"),
    ("bank_interest_lu", "6042 — of which from Luxembourg"),
    ("other_interest", "6043 — Other interest"),
    ("other_interest_lu", "6044 — of which from Luxembourg"),
    ("other_income", "6050 — Other income"),
    ("rental_income", "6051 — Rental income"),
    ("commission_income", "6052 — Commission income"),
    ("other_income_other", "6053 — Other"),
    ("charges", "6060 — Charges"),
    ("advisory_mgmt_fees", "6061 — Advisory and/or management commissions and/or fees"),
    ("depositary_fees", "6062 — Depositary commissions and/or fees"),
    ("admin_fees", "6063 — UCI administration fees"),
    ("audit_fees", "6064 — Audit and inspection fees"),
    ("distribution_fees", "6065 — Distribution fees"),
    ("other_admin_expenses", "6066 — Other administrative expenses"),
    ("subscription_tax", "6067 — Subscription tax"),
    ("other_taxes", "6068 — Other tax"),
    ("performance_fees", "6070 — Performance fees"),
    ("other_expenses", "6080 — Other expenses"),
    ("interest_paid", "6081 — Interest paid"),
    ("personnel_expenses", "6082 — Personnel expenses"),
    ("rental_expenses", "6083 — Rental expenses"),
    ("amortisation", "6084 — Amortisation"),
    ("other_expenses_other", "6085 — Other"),
    ("net_realised", "6090 — Net realised profit or loss on investments and foreign exchange positions"),
    ("net_unrealised", "6100 — Net unrealised profit or loss on investments and foreign exchange positions"),
]

# Short, single-line labels for the manual-entry grid — same rationale as
# SHORT_LABELS above; the full Annex 2 wording is always in the tooltip.
INCOME_SHORT_LABELS: dict[str, str] = {
    "dividends_shares": "6010 — Dividends on shares",
    "dividends_shares_lu": "6011 — of which from Luxembourg",
    "dividends_uci_units": "6020 — Dividends on UCI units/shares",
    "dividends_uci_units_lu": "6021 — of which from Luxembourg",
    "interest_bonds": "6030 — Interest on bonds/debt securities",
    "interest_bonds_lu": "6031 — of which from Luxembourg",
    "interest_loans_banks": "6040 — Interest on loans & accounts",
    "bank_interest": "6041 — Bank interest",
    "bank_interest_lu": "6042 — of which from Luxembourg",
    "other_interest": "6043 — Other interest",
    "other_interest_lu": "6044 — of which from Luxembourg",
    "other_income": "6050 — Other income",
    "rental_income": "6051 — Rental income",
    "commission_income": "6052 — Commission income",
    "other_income_other": "6053 — Other",
    "charges": "6060 — Charges",
    "advisory_mgmt_fees": "6061 — Advisory/management fees",
    "depositary_fees": "6062 — Depositary fees",
    "admin_fees": "6063 — UCI administration fees",
    "audit_fees": "6064 — Audit/inspection fees",
    "distribution_fees": "6065 — Distribution fees",
    "other_admin_expenses": "6066 — Other admin expenses",
    "subscription_tax": "6067 — Subscription tax",
    "other_taxes": "6068 — Other tax",
    "performance_fees": "6070 — Performance fees",
    "other_expenses": "6080 — Other expenses",
    "interest_paid": "6081 — Interest paid",
    "personnel_expenses": "6082 — Personnel expenses",
    "rental_expenses": "6083 — Rental expenses",
    "amortisation": "6084 — Amortisation",
    "other_expenses_other": "6085 — Other",
    "net_realised": "6090 — Net realised gains/losses",
    "net_unrealised": "6100 — Net unrealised gains/losses",
}
