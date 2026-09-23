"""Generation of the ingestion Excel template and reading/validation of the
files uploaded by ManCos (data per fund/sub-fund/share class).

The Excel template is the primary input channel of the mockup (see slide
"INGESTION — Excel templates" of the opportunity deck): each row describes
one share class and feeds directly into the PRIIPs calculation engine and
KID generation.
"""

from __future__ import annotations

import io
import re
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from domain.shared.sample_data import FUND_COLUMNS, SAMPLE_FUNDS

REQUIRED_COLUMNS = [
    "fund_id", "manco", "fund_name", "share_class", "isin", "regime", "category",
    "currency", "rhp_years", "asset_class", "mu_pct", "sigma_pct",
    "mgmt_fee_pct", "other_ongoing_pct", "transaction_cost_pct",
]

COLUMN_LABELS = {
    "fund_id": "Internal identifier", "manco": "ManCo", "promoter": "Promoter",
    "fund_name": "Fund name", "sub_fund": "Sub-fund", "share_class": "Share class",
    "isin": "ISIN code", "legal_form": "Legal form", "regime": "Regime (UCITS / AIF)",
    "category": "PRIIPs category (1-4)", "currency": "Currency", "launch_date": "Launch date",
    "rhp_years": "Recommended holding period (years)", "asset_class": "Asset class",
    "benchmark": "Benchmark", "mu_pct": "Expected annual return (%)",
    "sigma_pct": "Annualised volatility (%)", "stress_sigma_mult": "Stress volatility multiplier",
    "credit_quality_step": "Credit quality step (0-6)", "duration_years": "Duration (years)",
    "entry_cost_pct": "Entry cost (%)", "exit_cost_pct": "Exit cost (%)",
    "mgmt_fee_pct": "Management fee (%)", "other_ongoing_pct": "Other ongoing costs (%)",
    "transaction_cost_pct": "Transaction costs (%)", "perf_fee_pct": "Performance fee (%)",
    "objective_text": "Investment objective (text)", "investor_profile_text": "Target investor profile (text)",
}

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")


def build_ingestion_template_workbook(include_examples: bool = True) -> bytes:
    wb = openpyxl.Workbook()

    notice = wb.active
    notice.title = "Notes"
    notice["A1"] = "Ingestion template — PRIIPs KID Reporting"
    notice["A1"].font = Font(size=14, bold=True, color="0A2647")
    notice["A2"] = "One row = one share class. Fill in the 'Funds' sheet. Required columns are highlighted in the header."
    notice["A2"].alignment = Alignment(wrap_text=True)
    notice.column_dimensions["A"].width = 40
    notice.column_dimensions["B"].width = 90
    notice["A4"] = "Column"
    notice["B4"] = "Description"
    for cell in ("A4", "B4"):
        notice[cell].font = Font(bold=True)
        notice[cell].fill = PatternFill("solid", fgColor="0A2647")
        notice[cell].font = Font(bold=True, color="FFFFFF")
    r = 5
    for col in FUND_COLUMNS:
        notice[f"A{r}"] = col
        notice[f"B{r}"] = COLUMN_LABELS.get(col, col)
        r += 1

    ws = wb.create_sheet("Funds")
    header_fill_required = PatternFill("solid", fgColor="0A2647")
    header_fill_optional = PatternFill("solid", fgColor="5B6B82")
    for idx, col in enumerate(FUND_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill_required if col in REQUIRED_COLUMNS else header_fill_optional
        ws.column_dimensions[cell.column_letter].width = max(14, len(col) + 2)
        cell.comment = None

    ws.freeze_panes = "A2"

    if include_examples:
        for r_idx, fund in enumerate(SAMPLE_FUNDS[:3], start=2):
            for c_idx, col in enumerate(FUND_COLUMNS, start=1):
                ws.cell(row=r_idx, column=c_idx, value=fund.get(col))

    regime_col = FUND_COLUMNS.index("regime") + 1
    category_col = FUND_COLUMNS.index("category") + 1
    dv_regime = DataValidation(type="list", formula1='"UCITS,AIF"', allow_blank=True)
    dv_category = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
    ws.add_data_validation(dv_regime)
    ws.add_data_validation(dv_category)
    dv_regime.add(f"{ws.cell(row=2, column=regime_col).coordinate}:{ws.cell(row=500, column=regime_col).coordinate}")
    dv_category.add(f"{ws.cell(row=2, column=category_col).coordinate}:{ws.cell(row=500, column=category_col).coordinate}")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def parse_ingestion_workbook(file_bytes: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    errors: list[str] = []
    funds: list[dict[str, Any]] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        return [], [f"Unable to read the Excel file: {exc}"]

    if "Funds" not in wb.sheetnames:
        return [], ["Sheet 'Funds' not found. Please use the template provided by the application."]

    ws = wb["Funds"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        errors.append("Missing required columns: " + ", ".join(missing))
        return [], errors

    col_index = {name: idx for idx, name in enumerate(header)}

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        record: dict[str, Any] = {}
        for col in FUND_COLUMNS:
            idx = col_index.get(col)
            record[col] = row[idx] if idx is not None and idx < len(row) else None

        row_errors = _validate_row(record, row_num)
        if row_errors:
            errors.extend(row_errors)
            continue
        funds.append(_coerce_types(record))

    if not funds and not errors:
        errors.append("No data row found in the 'Funds' sheet.")

    return funds, errors


def _validate_row(record: dict[str, Any], row_num: int) -> list[str]:
    errs = []
    for col in REQUIRED_COLUMNS:
        if record.get(col) in (None, ""):
            errs.append(f"Row {row_num}: required field '{COLUMN_LABELS.get(col, col)}' is missing.")

    isin = record.get("isin")
    if isin and not ISIN_RE.match(str(isin).strip().upper()):
        errs.append(f"Row {row_num}: ISIN code '{isin}' is invalid (expected format: 2-letter country code + 10 characters).")

    category = record.get("category")
    if category not in (None, "") and str(category) not in {"1", "2", "3", "4"}:
        errs.append(f"Row {row_num}: PRIIPs category '{category}' is invalid (expected 1 to 4).")

    currency = record.get("currency")
    if currency and len(str(currency).strip()) != 3:
        errs.append(f"Row {row_num}: currency '{currency}' is invalid (expected a 3-letter ISO code).")

    for pct_col in ("mu_pct", "sigma_pct", "mgmt_fee_pct", "other_ongoing_pct", "transaction_cost_pct"):
        val = record.get(pct_col)
        if val not in (None, ""):
            try:
                float(val)
            except (TypeError, ValueError):
                errs.append(f"Row {row_num}: a numeric value was expected for '{COLUMN_LABELS.get(pct_col, pct_col)}' (received '{val}').")

    return errs


def _coerce_types(record: dict[str, Any]) -> dict[str, Any]:
    numeric_fields = [
        "rhp_years", "mu_pct", "sigma_pct", "stress_sigma_mult", "credit_quality_step",
        "duration_years", "entry_cost_pct", "exit_cost_pct", "mgmt_fee_pct",
        "other_ongoing_pct", "transaction_cost_pct", "perf_fee_pct", "category",
    ]
    out = dict(record)
    for f in numeric_fields:
        if out.get(f) not in (None, ""):
            try:
                out[f] = float(out[f])
            except (TypeError, ValueError):
                pass
    if out.get("category") is not None:
        out["category"] = int(out["category"])
    if out.get("rhp_years") is not None:
        out["rhp_years"] = int(out["rhp_years"])
    if out.get("isin"):
        out["isin"] = str(out["isin"]).strip().upper()
    if out.get("currency"):
        out["currency"] = str(out["currency"]).strip().upper()
    for missing_default, val in (("stress_sigma_mult", 1.5), ("credit_quality_step", 0), ("duration_years", 0.0), ("perf_fee_pct", 0.0), ("entry_cost_pct", 0.0), ("exit_cost_pct", 0.0)):
        if out.get(missing_default) in (None, ""):
            out[missing_default] = val
    return out
